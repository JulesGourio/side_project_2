"""utils.py — Docling-based parsing & chunking engine (PySpark / Databricks).

Responsibilities (stateless, runs on both driver and workers):
  - Centralised configuration injected from the driver via configure().
  - Docling DocumentConverter (PDF, DOCX, PPTX, XLSX, HTML, MD, TXT) with
    converters cached per worker process.
  - Passage splitting of the parsed markdown: chunking.py (section-aware, overlap, size caps).
  - Legacy-format pre-conversion (.doc/.rtf/.odt/.ods/.xls) and an XLSX fallback.
  - PySpark pandas UDFs for chunking and token counting.

configure() runs on the driver. Workers restore CONFIG via ensure_config()
(broadcast → spark.conf JSON → Volume JSON) or fall back to _DEFAULTS.
"""

import io
import os
import zipfile

# Must be set before any `transformers` import.
os.environ.setdefault("USE_TF", "0")
os.environ.setdefault("TRANSFORMERS_NO_TF", "1")
os.environ.setdefault("TF_CPP_MIN_LOG_LEVEL", "3")
os.environ.setdefault("CUDA_MODULE_LOADING", "LAZY")

import re
import time
import tempfile
import subprocess
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

import chunking  # shipped with addPyFile, like this module

from pyspark.sql import types as T
from pyspark.sql.functions import pandas_udf

# ---------------------------------------------------------------------------
# Logging — configure only our module logger, never the root logger
# (root=INFO leaks py4j/pyspark internal chatter to stdout).
# ---------------------------------------------------------------------------
logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)
if not logger.handlers:
    _handler = logging.StreamHandler()
    _handler.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s", datefmt="%H:%M:%S"))
    logger.addHandler(_handler)
    logger.propagate = False

for _noisy in ("py4j", "py4j.clientserver", "py4j.java_gateway", "pyspark", "docling", "transformers"):
    logging.getLogger(_noisy).setLevel(logging.WARNING)


# ===========================================================================
# Configuration
# ===========================================================================
# _DEFAULTS is the worker fallback; the notebook injects the real config via configure().
# ---------------------------------------------------------------------------
_DEFAULT_VOLUME_BASE_PATH = os.environ.get("PARSING_VOLUME_BASE_PATH")

_DEFAULTS: Dict[str, Any] = {
    # None (not a hardcoded path) if unset, so a bad fallback fails loudly downstream.
    "OFFLINE_MODELS_DIR": os.environ.get("PARSING_OFFLINE_MODELS"),
    "ANTIWORD_BIN": None,
    "ANTIWORD_SHARE_DIR": None,
    "VOLUME_BASE_PATH": _DEFAULT_VOLUME_BASE_PATH,
    # Engine knobs (balanced defaults)
    "USE_GPU": False,
    "DO_OCR": False,                      # True only if scanned PDFs are expected
    "GPU_OCR_FALLBACK": False,             # Docling's own do_ocr=True re-parse (EasyOCR, GPU) -- distinct from the LLM-OCR fallback
    "TABLE_STRUCTURE_MODE": "accurate",   # "accurate" | "fast"
    "GENERATE_PICTURE_IMAGES": True,      # set False when not describing images
    "IMAGE_SCALE": 0.75,
    # Image extraction
    "MIN_AREA_RATIO": 0.03,
    "MAX_REPEAT": 3,
    # Tokenizer
    "TOKENIZER_ENCODING": "cl100k_base",
    "USE_TIKTOKEN": False,
    "CHARS_PER_TOKEN": 3.5,
    # Chunking
    "MIN_CHUNK_TOKENS": 250,
    "TARGET_CHUNK_TOKENS": 500,
    "MAX_CHUNK_TOKENS": 1000,
    "MAX_CHUNK_CHARS": 4000,
    "CHUNK_OVERLAP_RATIO": 0.12,
    # LLM (image description)
    "LLM_MODEL_ENDPOINT": "databricks-gpt-5-nano",
    "LLM_MAX_TOKENS": 1024,
    "LLM_TEMPERATURE": 0.1,
    "LLM_MAX_RETRIES": 5,
    "LLM_MAX_CONCURRENT": 5,
    # Worker-config fallback file (must match VOLUME_BASE_PATH/_parsing_config.json)
    "WORKER_CONFIG_JSON": os.path.join(_DEFAULT_VOLUME_BASE_PATH, "_parsing_config.json"),
}

CONFIG: Dict[str, Any] = {}


def configure(**kwargs) -> None:
    """Set run configuration from the notebook. Call on the driver before parsing."""
    CONFIG.update(kwargs)
    logger.info("utils.configure() called with %d keys", len(kwargs))


# --- Multi-worker CONFIG propagation -------------------------------------
_CONFIG_BROADCAST = None


def broadcast_config(spark_session) -> None:
    """Broadcast CONFIG to all workers. Call once on the driver after configure()."""
    global _CONFIG_BROADCAST
    _CONFIG_BROADCAST = spark_session.sparkContext.broadcast(dict(CONFIG))
    logger.info("CONFIG broadcast (%d keys)", len(CONFIG))


def write_worker_config(spark_session) -> str:
    """Persist CONFIG to a Volume JSON + spark.conf so USER_ISOLATION workers can
    restore it (spark.conf is not readable inside pandas UDFs on shared clusters).
    Returns the JSON path written."""
    import json
    path = _cfg("WORKER_CONFIG_JSON", fallback=os.path.join(_cfg("VOLUME_BASE_PATH"), "_parsing_config.json"))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w") as f:
        json.dump(CONFIG, f)
    try:
        spark_session.conf.set("app.parsing.config_json", json.dumps(CONFIG))
    except Exception:
        pass
    logger.info("Worker config written to %s (%d keys)", path, len(CONFIG))
    return path


def ensure_config() -> None:
    """Restore CONFIG on a worker (idempotent): broadcast → spark.conf → Volume JSON."""
    global _CONFIG_BROADCAST
    if CONFIG:
        return
    if _CONFIG_BROADCAST is not None:
        CONFIG.update(_CONFIG_BROADCAST.value)
        return
    import json
    try:
        from pyspark.sql import SparkSession
        _spark = SparkSession.getActiveSession()
        if _spark:
            cfg_json = _spark.conf.get("app.parsing.config_json", "")
            if cfg_json:
                CONFIG.update(json.loads(cfg_json))
                return
    except Exception:
        pass
    try:
        # A cluster env var, visible in every process unlike CONFIG.
        volume_base = os.environ.get("PARSING_VOLUME_BASE_PATH")
        vol_cfg = os.path.join(volume_base, "_parsing_config.json") if volume_base \
            else _DEFAULTS["WORKER_CONFIG_JSON"]
        if os.path.isfile(vol_cfg):
            with open(vol_cfg) as f:
                CONFIG.update(json.loads(f.read()))
    except Exception:
        pass  # Falls back to _DEFAULTS via _cfg()


def _cfg(key: str, fallback=None):
    """Read a config value: CONFIG → _DEFAULTS → fallback → error."""
    if key in CONFIG:
        return CONFIG[key]
    if key in _DEFAULTS and _DEFAULTS[key] is not None:
        return _DEFAULTS[key]
    if fallback is not None:
        return fallback
    if key in _DEFAULTS:           # default is legitimately None
        return None
    raise RuntimeError(f"Configuration key '{key}' not set and no default exists.")


# ===========================================================================
# Non-configurable internal constants
# ===========================================================================
TEXTLIKE_EXTENSIONS = {"csv", "json", "tsv"}
CONVERTED_EXTS = {".doc", ".rtf", ".odt", ".ods", ".xls"}

_FORMAT_MAP_NAMES: Dict[str, str] = {
    ".pdf":  "PDF", ".docx": "DOCX", ".docm": "DOCX", ".pptx": "PPTX",
    ".xlsx": "XLSX", ".xlsm": "XLSX", ".xlsb": "XLSX", ".html": "HTML",
    ".htm":  "HTML", ".xml":  "HTML", ".md":   "MD", ".txt":  "MD",
}

_TIKTOKEN_ENCODER = None

# ---------------------------------------------------------------------------
# Spark schemas
# ---------------------------------------------------------------------------
CHUNK_SCHEMA = T.ArrayType(T.StructType([
    T.StructField("chunk_index", T.IntegerType(), True),
    T.StructField("chunk_text", T.StringType(), True),
    T.StructField("chunk_char_count", T.IntegerType(), True),
    T.StructField("chunk_token_count", T.IntegerType(), True),
    T.StructField("chunk_content_type", T.StringType(), True),
    T.StructField("metadata", T.MapType(T.StringType(), T.StringType()), True),
]))


# ===========================================================================
# Docling bootstrap (lazy, cached per worker)
# ===========================================================================
def _patch_worker_env():
    """Set offline / thread-limit env vars on Spark workers."""
    os.environ.setdefault("USER", "spark_worker")
    os.environ["USE_TF"] = "0"
    os.environ["TRANSFORMERS_NO_TF"] = "1"
    os.environ["TF_CPP_MIN_LOG_LEVEL"] = "3"
    os.environ["CUDA_MODULE_LOADING"] = "LAZY"
    os.environ["HF_HUB_OFFLINE"] = "1"
    os.environ["TRANSFORMERS_OFFLINE"] = "1"
    os.environ["HF_HUB_DISABLE_TELEMETRY"] = "1"
    n_threads = "2" if _cfg("USE_GPU", fallback=False) else "1"
    for var in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS",
                "VECLIB_MAXIMUM_THREADS", "NUMEXPR_NUM_THREADS"):
        os.environ[var] = n_threads


def _import_docling():
    _patch_worker_env()
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import (
        PdfPipelineOptions, AcceleratorOptions, AcceleratorDevice,
    )

    mods = {
        "DocumentConverter": DocumentConverter, "PdfFormatOption": PdfFormatOption,
        "InputFormat": InputFormat, "PdfPipelineOptions": PdfPipelineOptions,
        "AcceleratorOptions": AcceleratorOptions, "AcceleratorDevice": AcceleratorDevice,
    }
    try:
        from docling.datamodel.pipeline_options import TableFormerMode
        mods["TableFormerMode"] = TableFormerMode
    except Exception:
        mods["TableFormerMode"] = None

    if _cfg("USE_GPU", fallback=False):
        try:
            import torch
            if torch.cuda.is_available():
                props = torch.cuda.get_device_properties(0)
                total = props.total_mem / (1024 ** 3) if hasattr(props, "total_mem") else 0.0
                logger.info("GPU: %s | %.1f GB | CUDA %s",
                            torch.cuda.get_device_name(0), total, torch.version.cuda)
            else:
                logger.warning("USE_GPU=True but CUDA unavailable — falling back to CPU")
        except Exception as e:
            logger.warning("GPU diagnostics failed: %s", str(e)[:200])
    return mods


_DOCLING_MODULES = None
_CONVERTER_CACHE: Dict = {}


def _get_docling():
    global _DOCLING_MODULES
    if _DOCLING_MODULES is None:
        _DOCLING_MODULES = _import_docling()
    return _DOCLING_MODULES


def _get_input_format(ext: str):
    return getattr(_get_docling()["InputFormat"], _FORMAT_MAP_NAMES.get(ext, ""), None)


def _get_converter(input_format, do_ocr: Optional[bool] = None, timings: Optional[dict] = None):
    """Build or retrieve a cached DocumentConverter for the given format.

    The cache key includes GPU mode, table mode, OCR and picture-image flags so
    that overriding any knob (e.g. an OCR retry) yields a distinct converter.
    """
    dl = _get_docling()
    use_gpu = _cfg("USE_GPU", fallback=False)
    table_mode = str(_cfg("TABLE_STRUCTURE_MODE", fallback="accurate")).lower()
    gen_images = bool(_cfg("GENERATE_PICTURE_IMAGES", fallback=True))
    do_ocr = bool(_cfg("DO_OCR", fallback=False)) if do_ocr is None else do_ocr

    fmt_key = f"{input_format}|gpu={use_gpu}|tbl={table_mode}|ocr={do_ocr}|img={gen_images}"
    if fmt_key in _CONVERTER_CACHE:
        return _CONVERTER_CACHE[fmt_key]

    t0 = time.perf_counter()
    opts = dl["PdfPipelineOptions"]()
    opts.artifacts_path = Path(_cfg("OFFLINE_MODELS_DIR"))
    opts.do_ocr = do_ocr
    opts.do_table_structure = True
    opts.generate_picture_images = gen_images
    opts.images_scale = _cfg("IMAGE_SCALE", fallback=0.75)
    if dl["TableFormerMode"] is not None:
        try:
            opts.table_structure_options.mode = (
                dl["TableFormerMode"].FAST if table_mode == "fast" else dl["TableFormerMode"].ACCURATE
            )
        except Exception as e:
            logger.warning("Could not set TableFormerMode=%s: %s", table_mode, str(e)[:120])

    if use_gpu:
        opts.accelerator_options = dl["AcceleratorOptions"](num_threads=4, device=dl["AcceleratorDevice"].CUDA)
    else:
        opts.accelerator_options = dl["AcceleratorOptions"](num_threads=1, device=dl["AcceleratorDevice"].CPU)

    format_options = {}
    if input_format == dl["InputFormat"].PDF:
        backend = None
        try:
            from docling.backend.pypdfium2_backend import PyPdfiumDocumentBackend as backend  # faster text layer
        except ImportError:
            try:
                from docling.backend.docling_parse_backend import DoclingParseDocumentBackend as backend
            except ImportError:
                backend = None
        format_options[dl["InputFormat"].PDF] = (
            dl["PdfFormatOption"](pipeline_options=opts, backend=backend) if backend
            else dl["PdfFormatOption"](pipeline_options=opts)
        )

    converter = dl["DocumentConverter"](format_options=format_options)
    if timings is not None:
        timings["docling_load_seconds"] = float(time.perf_counter() - t0)
    _CONVERTER_CACHE[fmt_key] = converter
    return converter


# ===========================================================================
# Text helpers
# ===========================================================================
def safe_decode(raw_bytes: bytes, encodings: Optional[List[str]] = None) -> str:
    if not raw_bytes:
        return ""
    for enc in (encodings or ["utf-8", "utf-8-sig", "cp1252", "latin-1"]):
        try:
            return raw_bytes.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw_bytes.decode("utf-8", errors="ignore")


def normalize_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def get_tiktoken_encoder():
    global _TIKTOKEN_ENCODER
    if _TIKTOKEN_ENCODER is None:
        import tiktoken
        _TIKTOKEN_ENCODER = tiktoken.get_encoding(_cfg("TOKENIZER_ENCODING", fallback="cl100k_base"))
    return _TIKTOKEN_ENCODER


def count_tokens(text: str) -> int:
    """Token count via tiktoken (USE_TIKTOKEN=True) or a chars/token estimate."""
    if not text:
        return 0
    if _cfg("USE_TIKTOKEN", fallback=False):
        try:
            return len(get_tiktoken_encoder().encode(text))
        except Exception:
            pass
    return max(1, int(len(text) / _cfg("CHARS_PER_TOKEN", fallback=3.5)))


def _write_tmp_file(content: bytes, suffix: str) -> Path:
    tmp = tempfile.NamedTemporaryFile(suffix=suffix, delete=False, prefix="docling_parse_")
    tmp.write(content)
    tmp.close()
    return Path(tmp.name)


# ===========================================================================
# Legacy-format pre-conversion
# ===========================================================================
def _convert_doc_antiword(doc_path: Path) -> Tuple[Optional[Path], Optional[str]]:
    """Convert legacy .doc to Markdown via the Antiword binary (if configured).

    Returns (path, None) on success or (None, reason) on failure.
    """
    antiword_bin = _cfg("ANTIWORD_BIN")
    antiword_share = _cfg("ANTIWORD_SHARE_DIR")
    if not antiword_bin or not os.path.exists(antiword_bin):
        return None, f"binary not found at {antiword_bin}"
    if not os.access(antiword_bin, os.X_OK):
        try:
            os.chmod(antiword_bin, 0o755)
        except OSError as e:
            return None, f"binary not executable and chmod failed: {str(e)[:200]}"
    try:
        env = os.environ.copy()
        if antiword_share:
            env["ANTIWORDHOME"] = antiword_share
        result = subprocess.run(
            [antiword_bin, "-m", "UTF-8.txt", "-w", "0", str(doc_path)],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            check=False, env=env, timeout=120,
        )
        if result.returncode != 0 or not result.stdout.strip():
            return None, f"exit={result.returncode} stderr={(result.stderr or '')[:300]}"
        tmp = tempfile.NamedTemporaryFile(mode="w", suffix=".md", encoding="utf-8",
                                          delete=False, prefix=f"{doc_path.stem}_antiword_")
        tmp.write(result.stdout)
        tmp.close()
        return Path(tmp.name), None
    except subprocess.TimeoutExpired:
        return None, "timeout after 120s"
    except Exception as e:
        return None, f"raised {type(e).__name__}: {str(e)[:300]}"


def _convert_rtf(content: bytes) -> Tuple[Optional[str], Optional[str]]:
    """Returns (text, None) on success or (None, error_detail) on failure —
    the real cause is preserved instead of a generic string."""
    try:
        from striprtf.striprtf import rtf_to_text
        text = rtf_to_text(content.decode("utf-8", errors="replace"))
        return (text, None) if text and text.strip() else (None, "striprtf returned no text")
    except Exception as e:
        return None, _exc_detail(e)


def _convert_odt(content: bytes, suffix: str) -> Tuple[Optional[str], Optional[str]]:
    """Returns (text, None) on success or (None, error_detail) on failure —
    the real cause is preserved instead of a generic string."""
    try:
        from odf.opendocument import load as odf_load
        from odf.text import P
        from odf.table import Table, TableRow, TableCell
    except ImportError as e:
        return None, f"odfpy import failed: {e}"
    tmp_path = _write_tmp_file(content, suffix)
    try:
        doc = odf_load(str(tmp_path))
        lines = []
        if suffix == ".ods":
            for sheet in doc.spreadsheet.getElementsByType(Table):
                lines.append(f"## {sheet.getAttribute('name') or 'Sheet'}")
                for row in sheet.getElementsByType(TableRow):
                    cells = [" ".join(str(p) for p in cell.getElementsByType(P)).strip()
                             for cell in row.getElementsByType(TableCell)]
                    if any(cells):
                        lines.append("\t".join(cells))
        else:
            for p in doc.text.getElementsByType(P):
                lines.append(str(p))
        text = "\n".join(lines)
        return (text, None) if text.strip() else (None, "ODF document yielded no text")
    except Exception as e:
        return None, _exc_detail(e)
    finally:
        try:
            os.unlink(str(tmp_path))
        except OSError:
            pass


def _convert_xls(content: bytes) -> Tuple[Optional[str], Optional[str]]:
    """Returns (text, None) on success or (None, error_detail) on failure —
    the real cause is preserved instead of a generic string."""
    try:
        import xlrd
    except ImportError as e:
        return None, f"xlrd import failed: {e}"
    tmp_path = _write_tmp_file(content, ".xls")
    try:
        wb = xlrd.open_workbook(str(tmp_path))
        lines = []
        for sheet in wb.sheets():
            rows = chunking.sheet_lines([[sheet.cell_value(rx, cx) for cx in range(sheet.ncols)]
                                         for rx in range(sheet.nrows)])
            if rows:
                lines.append(f"## {sheet.name}\n" + "\n".join(rows))
        text = "\n\n".join(lines)
        return (text, None) if text.strip() else (None, "XLS workbook yielded no text")
    except Exception as e:
        return None, _exc_detail(e)
    finally:
        try:
            os.unlink(str(tmp_path))
        except OSError:
            pass


# ===========================================================================
# Core parsing
# ===========================================================================
def parse_with_docling(content: bytes, extension: str, timings: Optional[dict] = None) -> Dict[str, Any]:
    """Route raw bytes to the right parser by extension. Returns a dict with
    text / parser_error / parser_strategy / parse_time_seconds / _docling_doc.
    If ``timings`` is provided it is filled with per-step durations."""
    ext = f".{extension.lower().lstrip('.')}"
    start = time.perf_counter()

    if ext.lstrip(".") in TEXTLIKE_EXTENSIONS:
        return _parse_textlike(content, ext, start)
    if ext == ".doc":
        return _parse_doc_antiword(content, start)
    if ext == ".rtf":
        return _parse_rtf(content, start)
    if ext in {".odt", ".ods"}:
        return _parse_odf(content, ext, start)
    if ext == ".xls":
        return _parse_xls(content, start)

    # Fast-fail for invalid DOCX/DOCM (otherwise Docling can hang ~45s).
    if ext in {".docx", ".docm"}:
        import zipfile as _zf
        try:
            with _zf.ZipFile(io.BytesIO(content)) as zfh:
                if "word/document.xml" not in zfh.namelist():
                    return _err(start, ext, "Invalid DOCX (no word/document.xml)")
        except Exception:
            return _err(start, ext, "Invalid DOCX (not a valid ZIP)")

    t_imp = time.perf_counter()
    input_format = _get_input_format(ext)
    if timings is not None:
        timings["docling_import_seconds"] = float(time.perf_counter() - t_imp)
    if input_format is None:
        return _err(start, ext, f"Unsupported format: {ext}", strategy=f"unsupported:{ext.lstrip('.')}")

    return _parse_via_docling(content, ext, input_format, start, timings)


def _exc_detail(e: Exception, limit: int = 500) -> str:
    """Walks __cause__/__context__ instead of str(e) alone — Docling/PyTorch
    exceptions frequently wrap a deeper root cause this way."""
    parts = [str(e)]
    cause = e.__cause__ or e.__context__
    seen = {id(e)}
    while cause is not None and id(cause) not in seen:
        parts.append(str(cause))
        seen.add(id(cause))
        cause = cause.__cause__ or cause.__context__
    return " | caused by: ".join(parts)[:limit]


def _err(start, ext, msg, strategy=None) -> Dict[str, Any]:
    return {
        "text": "", "parser_error": msg,
        "parser_strategy": strategy or f"docling_error:{ext.lstrip('.')}",
        "parse_time_seconds": float(time.perf_counter() - start), "_docling_doc": None,
    }


def _export_markdown(doc) -> str:
    try:
        from docling_core.types.doc.document import ImageRefMode
        return doc.export_to_markdown(image_mode=ImageRefMode.PLACEHOLDER)
    except (ImportError, AttributeError, TypeError):
        return doc.export_to_markdown()


# ---------------------------------------------------------------------------
# Pre-clean OOXML zips (.docx/.docm/.pptx/.pptm) before the size gate.
# ---------------------------------------------------------------------------
_OOXML_EMBEDDING_MARKER = "/embeddings/"
_MEDIA_JUNK_EXTS = {
    ".mp4", ".avi", ".mov", ".wmv", ".mpg", ".mpeg", ".m4v", ".flv", ".mkv", ".asf",
    ".mp3", ".wav", ".m4a", ".wma", ".aac",
}
_MEDIA_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}
_RECOMPRESS_MAX_DIM = 1600
_RECOMPRESS_JPEG_QUALITY = 78


def _recompress_image_bytes(data: bytes) -> bytes:
    """Downscale/recompress one embedded image; returns original bytes on any
    failure or if the result isn't actually smaller (e.g. already tiny)."""
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(data))
        img.load()
        fmt = img.format or "PNG"
        w, h = img.size
        if max(w, h) > _RECOMPRESS_MAX_DIM:
            scale = _RECOMPRESS_MAX_DIM / max(w, h)
            img = img.resize((max(1, int(w * scale)), max(1, int(h * scale))), Image.LANCZOS)
        out = io.BytesIO()
        save_kwargs = {"optimize": True}
        if fmt == "JPEG":
            save_kwargs["quality"] = _RECOMPRESS_JPEG_QUALITY
        img.save(out, format=fmt, **save_kwargs)
        new_bytes = out.getvalue()
        return new_bytes if len(new_bytes) < len(data) else data
    except Exception:
        return data


def strip_ooxml_bloat(content: bytes, ext: str) -> bytes:
    """Rebuild a .docx/.docm/.pptx/.pptm zip without embedded OLE/Office
    objects or embedded video/audio, and with large embedded images
    recompressed. Returns the ORIGINAL bytes unchanged on any error or if the
    file isn't a zip we recognise — this must never be the reason a file that
    would otherwise have parsed fine now fails."""
    if ext.lstrip(".").lower() not in {"docx", "docm", "pptx", "pptm"}:
        return content
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zin:
            names = set(zin.namelist())
            if "word/document.xml" not in names and "ppt/presentation.xml" not in names:
                return content
            out_buf = io.BytesIO()
            with zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    lower = item.filename.lower()
                    _, item_ext = os.path.splitext(lower)
                    if _OOXML_EMBEDDING_MARKER in lower:
                        continue  # embedded OLE/Office object — not needed for text
                    if "/media/" in lower and item_ext in _MEDIA_JUNK_EXTS:
                        continue  # embedded video/audio — not needed for text
                    data = zin.read(item.filename)
                    if "/media/" in lower and item_ext in _MEDIA_IMAGE_EXTS:
                        data = _recompress_image_bytes(data)
                    zout.writestr(item, data)
            return out_buf.getvalue()
    except Exception:
        return content


def _parse_via_docling(content, ext, input_format, start, timings) -> Dict[str, Any]:
    t_tmp = time.perf_counter()
    tmp_path = _write_tmp_file(content, ext)
    if timings is not None:
        timings["tmp_file_write_seconds"] = float(time.perf_counter() - t_tmp)
    try:
        converter = _get_converter(input_format, timings=timings)

        t0 = time.perf_counter()
        doc = converter.convert(str(tmp_path)).document
        if timings is not None:
            timings["docling_convert_seconds"] = float(time.perf_counter() - t0)

        t0 = time.perf_counter()
        md_text = normalize_text(_export_markdown(doc))
        if timings is not None:
            timings["markdown_export_seconds"] = float(time.perf_counter() - t0)

        # OCR fallback: re-parse scanned PDFs that yielded (near) no text.
        ocr_failure_detail = None
        if (not md_text or len(md_text) < 20) and ext == ".pdf" and _cfg("GPU_OCR_FALLBACK", fallback=False):
            t_ocr = time.perf_counter()
            try:
                ocr_conv = _get_converter(input_format, do_ocr=True)
                ocr_doc = ocr_conv.convert(str(tmp_path)).document
                ocr_md = normalize_text(_export_markdown(ocr_doc))
                if timings is not None:
                    timings["ocr_fallback_seconds"] = float(time.perf_counter() - t_ocr)
                if len(ocr_md) > len(md_text):
                    return {"text": ocr_md, "parser_error": None if ocr_md else "Empty after OCR",
                            "parser_strategy": "docling+ocr:pdf",
                            "parse_time_seconds": float(time.perf_counter() - start), "_docling_doc": ocr_doc}
            except Exception as e:
                if timings is not None:
                    timings["ocr_fallback_seconds"] = float(time.perf_counter() - t_ocr)
                ocr_failure_detail = _exc_detail(e, limit=300)
                logger.warning("OCR fallback failed: %s", ocr_failure_detail)

        if md_text:
            parser_error = None
        elif ocr_failure_detail:
            parser_error = f"Empty after Docling conversion; OCR fallback also failed: {ocr_failure_detail}"
        else:
            parser_error = "Empty after Docling conversion"
        return {"text": md_text, "parser_error": parser_error,
                "parser_strategy": f"docling:{ext.lstrip('.')}",
                "parse_time_seconds": float(time.perf_counter() - start), "_docling_doc": doc}
    except Exception as e:
        if ext in {".xlsx", ".xlsm", ".xlsb"}:
            fallback = _parse_xlsx_openpyxl(content, start, ext.lstrip("."))
            if fallback.get("text"):
                return fallback
        return _err(start, ext, _exc_detail(e))
    finally:
        try:
            os.unlink(str(tmp_path))
        except OSError:
            pass


def _parse_textlike(content: bytes, ext: str, start: float) -> Dict[str, Any]:
    decoded = safe_decode(content)
    if ext in {".csv", ".tsv"}:
        try:
            sep = "\t" if ext == ".tsv" else ","
            df = pd.read_csv(io.StringIO(decoded), sep=sep)
            decoded = "\n".join(
                f"- {', '.join(f'{h}: {row[h]}' for h in df.columns if pd.notna(row[h]))}"
                for _, row in df.iterrows()
            )
        except Exception:
            pass
    decoded = normalize_text(decoded)
    return {"text": decoded, "parser_error": None if decoded else "Empty",
            "parser_strategy": f"textlike:{ext.lstrip('.')}",
            "parse_time_seconds": float(round(time.perf_counter() - start, 2)), "_docling_doc": None}


def _parse_doc_antiword(content: bytes, start: float) -> Dict[str, Any]:
    """Parse .doc via antiword; reroute to DOCX if it is actually a renamed ZIP."""
    if content[:4] == b"PK\x03\x04":
        import zipfile as _zf
        try:
            with _zf.ZipFile(io.BytesIO(content)) as zfh:
                if "word/document.xml" in zfh.namelist():
                    return parse_with_docling(content, "docx")
        except Exception:
            pass

    tmp_doc = _write_tmp_file(content, ".doc")
    try:
        md_path, err_reason = _convert_doc_antiword(tmp_doc)
        if md_path is None:
            return {"text": "", "parser_error": f"antiword conversion failed: {err_reason}",
                    "parser_strategy": "antiword_failed:doc",
                    "parse_time_seconds": float(round(time.perf_counter() - start, 2)), "_docling_doc": None}
        try:
            with open(str(md_path), "r", encoding="utf-8") as f:
                text = normalize_text(f.read())
            strategy = "antiword:doc" if text else "antiword_empty:doc"
            return {"text": text, "parser_error": None, "parser_strategy": strategy,
                    "parse_time_seconds": float(round(time.perf_counter() - start, 2)), "_docling_doc": None}
        finally:
            try:
                os.unlink(str(md_path))
            except OSError:
                pass
    finally:
        try:
            os.unlink(str(tmp_doc))
        except OSError:
            pass


def _parse_rtf(content: bytes, start: float) -> Dict[str, Any]:
    text, err_reason = _convert_rtf(content)
    if not text or not text.strip():
        return {"text": "", "parser_error": f"RTF conversion failed: {err_reason}",
                "parser_strategy": "striprtf_failed:rtf",
                "parse_time_seconds": float(round(time.perf_counter() - start, 2)), "_docling_doc": None}
    tmp = tempfile.NamedTemporaryFile(mode="w", suffix=".md", encoding="utf-8", delete=False)
    tmp.write(text)
    tmp.close()
    try:
        doc = _get_converter(_get_docling()["InputFormat"].MD).convert(tmp.name).document
        return {"text": normalize_text(doc.export_to_markdown()), "parser_error": None,
                "parser_strategy": "striprtf+docling:rtf",
                "parse_time_seconds": float(round(time.perf_counter() - start, 2)), "_docling_doc": doc}
    except Exception:
        return {"text": normalize_text(text), "parser_error": None, "parser_strategy": "striprtf_raw:rtf",
                "parse_time_seconds": float(round(time.perf_counter() - start, 2)), "_docling_doc": None}
    finally:
        try:
            os.unlink(tmp.name)
        except OSError:
            pass


def _parse_odf(content: bytes, ext: str, start: float) -> Dict[str, Any]:
    raw_text, err_reason = _convert_odt(content, ext)
    text = normalize_text(raw_text or "")
    return {"text": text, "parser_error": None if text else f"ODF conversion failed: {err_reason}",
            "parser_strategy": f"odfpy:{ext.lstrip('.')}",
            "parse_time_seconds": float(round(time.perf_counter() - start, 2)), "_docling_doc": None}


def _parse_xls(content: bytes, start: float) -> Dict[str, Any]:
    raw_text, err_reason = _convert_xls(content)
    text = normalize_text(raw_text or "")
    return {"text": text, "parser_error": None if text else f"XLS conversion failed: {err_reason}",
            "parser_strategy": "xlrd:xls",
            "parse_time_seconds": float(round(time.perf_counter() - start, 2)), "_docling_doc": None}


def _parse_xlsb_pyxlsb(content: bytes, start: float) -> Dict[str, Any]:
    """.xlsb is a ZIP container like .xlsx, but its sheets are binary BIFF12
    records, not XML — openpyxl can't open it and the XML-raw fallback finds
    no xl/worksheets/sheetN.xml to read, so both silently yield empty text."""
    try:
        import pyxlsb
        all_text = []
        with pyxlsb.open_workbook(io.BytesIO(content)) as wb:
            for sheet_name in wb.sheets:
                with wb.get_sheet(sheet_name) as ws:
                    rows = chunking.sheet_lines([[c.v for c in row] for row in ws.rows()])
                    if rows:
                        all_text.append(f"## {sheet_name}\n" + "\n".join(rows))
        text = normalize_text("\n\n".join(all_text))
        return {"text": text, "parser_error": None if text else "xlsb empty (pyxlsb extraction)",
                "parser_strategy": "pyxlsb:xlsb" if text else "pyxlsb_empty:xlsb",
                "parse_time_seconds": float(time.perf_counter() - start), "_docling_doc": None}
    except Exception as e:
        return {"text": "", "parser_error": f"pyxlsb conversion failed: {_exc_detail(e)}",
                "parser_strategy": "pyxlsb_failed:xlsb",
                "parse_time_seconds": float(time.perf_counter() - start), "_docling_doc": None}


def _parse_xlsx_openpyxl(content: bytes, start: float, ext: str = "xlsx") -> Dict[str, Any]:
    """Fallback for .xlsx/.xlsm: openpyxl, then raw ZIP+XML extraction.
    .xlsb is routed to _parse_xlsb_pyxlsb instead — see its docstring."""
    if ext.lstrip(".") == "xlsb":
        return _parse_xlsb_pyxlsb(content, start)

    import zipfile as _zf
    import xml.etree.ElementTree as _ET

    # --- Attempt 1: openpyxl ---
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        all_text = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = chunking.sheet_lines(list(ws.iter_rows(values_only=True)))
            if rows:
                all_text.append(f"## {sheet_name}\n" + "\n".join(rows))
        wb.close()
        text = normalize_text("\n\n".join(all_text))
        if text:
            return {"text": text, "parser_error": None, "parser_strategy": "openpyxl:xlsx",
                    "parse_time_seconds": float(time.perf_counter() - start), "_docling_doc": None}
    except Exception:
        pass

    # --- Attempt 2: raw XML extraction ---
    try:
        zfh = _zf.ZipFile(io.BytesIO(content))
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        shared_strings = []
        if "xl/sharedStrings.xml" in zfh.namelist():
            try:
                ss_root = _ET.fromstring(zfh.read("xl/sharedStrings.xml"))
                for si in ss_root.findall(f"{{{ns}}}si"):
                    shared_strings.append("".join(t.text or "" for t in si.iter(f"{{{ns}}}t")))
            except _ET.ParseError:
                shared_strings = []

        all_text = []
        sheet_files = sorted(n for n in zfh.namelist()
                             if n.startswith("xl/worksheets/sheet") and n.endswith(".xml"))
        for sheet_idx, sheet_path in enumerate(sheet_files, 1):
            try:
                sheet_root = _ET.fromstring(zfh.read(sheet_path))
            except _ET.ParseError:
                continue
            rows_data = []
            for row_el in sheet_root.findall(f".//{{{ns}}}row"):
                row_vals = []
                for cell in row_el.findall(f"{{{ns}}}c"):
                    v_elem = cell.find(f"{{{ns}}}v")
                    if v_elem is not None and v_elem.text:
                        if cell.get("t", "") == "s" and v_elem.text.isdigit():
                            idx = int(v_elem.text)
                            row_vals.append(shared_strings[idx] if 0 <= idx < len(shared_strings) else "")
                        else:
                            row_vals.append(v_elem.text)
                    else:
                        row_vals.append("")
                if any(v.strip() for v in row_vals):
                    rows_data.append(row_vals)
            if rows_data:
                lines = chunking.sheet_lines(rows_data)
                if lines:
                    all_text.append(f"## Sheet {sheet_idx}\n" + "\n".join(lines))
        zfh.close()
        text = normalize_text("\n\n".join(all_text))
        return {"text": text, "parser_error": None if text else "xlsx empty (XML extraction)",
                "parser_strategy": "xml_raw:xlsx",
                "parse_time_seconds": float(time.perf_counter() - start), "_docling_doc": None}
    except Exception as e:
        return {"text": "", "parser_error": f"xlsx all fallbacks failed: {str(e)[:200]}",
                "parser_strategy": "xlsx_failed:xlsx",
                "parse_time_seconds": float(time.perf_counter() - start), "_docling_doc": None}


_INTRAQUAL_REF_URL_BASE = "https://intraqual.lat.corp/intraqual_prod/identification.aspx?ref="


def intraqual_ref_url(ref_col):
    """Document URL built directly from REF, no lookup table needed."""
    from pyspark.sql import functions as F
    return F.when(
        ref_col.isNotNull(),
        F.concat(F.lit(_INTRAQUAL_REF_URL_BASE), F.regexp_replace(ref_col, " ", "%20")),
    )


# ===========================================================================
# Chunking
# ===========================================================================
def source_prefixed_text(body_col, ref_col, titre_col, division_col, category_col,
                          doc_date_col=None, include_prefix=True, type_col=None):
    """chunk_text with an optional "[Source: ...]" provenance prefix (include_prefix).
    type_col (gd_typdoc label, e.g. "05 - Procédure - QP") adds a "Type:" field."""
    from pyspark.sql import functions as F
    if not include_prefix:
        return body_col
    date_part = (
        F.concat(F.lit(" | Date de diffusion: "), F.coalesce(F.date_format(doc_date_col, "yyyy-MM-dd"), F.lit("inconnue")))
        if doc_date_col is not None else F.lit("")
    )
    type_part = (F.concat(F.lit(" | Type: "), F.coalesce(type_col, F.lit("")))
                 if type_col is not None else F.lit(""))
    return F.concat(
        F.lit("[Source: "), F.coalesce(ref_col, F.lit("")),
        F.lit(" | Title: "), F.coalesce(titre_col, F.lit("")),
        type_part,
        F.lit(" | Division: "), F.coalesce(division_col, F.lit("")),
        F.lit(" | Category: "), F.coalesce(category_col, F.lit("")),
        date_part,
        F.lit("]\n\n"), body_col,
    )


def chunk_document(text: str, docling_doc=None) -> List[Dict[str, Any]]:
    """Passages of an already-parsed markdown document (``chunking.chunk_markdown``).

    ``docling_doc`` is accepted for the sandbox notebook's signature only: the pipeline keeps
    the markdown, not the Docling document, so passages always come from the markdown (the
    former Docling HybridChunker path never ran in the pipeline — audit 2026-10, P1).
    """
    text = normalize_text(text)
    if not text:
        return []
    return chunking.chunk_markdown(
        text,
        min_tokens=int(_cfg("MIN_CHUNK_TOKENS", fallback=250)),
        target_tokens=int(_cfg("TARGET_CHUNK_TOKENS", fallback=500)),
        max_tokens=int(_cfg("MAX_CHUNK_TOKENS", fallback=1000)),
        max_chars=int(_cfg("MAX_CHUNK_CHARS", fallback=4000)),
        overlap_ratio=float(_cfg("CHUNK_OVERLAP_RATIO", fallback=0.12)),
        chars_per_token=float(_cfg("CHARS_PER_TOKEN", fallback=3.5)),
        count_tokens=count_tokens,
    )


# ===========================================================================
# PySpark pandas UDFs
# ===========================================================================
@pandas_udf(CHUNK_SCHEMA)
def build_chunks_udf(text_series: pd.Series) -> pd.Series:
    """Chunk already-parsed markdown text (chunking.chunk_markdown)."""
    ensure_config()
    return pd.Series([
        chunk_document(str(t), docling_doc=None) if t else [] for t in text_series
    ])


@pandas_udf(T.StringType())
def language_udf(text_series: pd.Series, ref_series: pd.Series) -> pd.Series:
    """Document language (chunking.detect_language): REF suffix, else stop words."""
    return pd.Series([chunking.detect_language(str(t or ""), str(r or ""))
                      for t, r in zip(text_series, ref_series)])


@pandas_udf(T.IntegerType())
def token_count_udf(text_series: pd.Series) -> pd.Series:
    ensure_config()
    return text_series.apply(
        lambda x: count_tokens(str(x)) if pd.notna(x) and str(x).strip() else 0
    ).astype("int32")
