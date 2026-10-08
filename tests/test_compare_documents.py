"""End-to-end extraction + diff on generated PDF and DOCX documents.

Each test builds two revisions of a small quality procedure with ONE known
difference and checks what reaches the LLM: the real change, and nothing else.
Documents are generated on the fly (PyMuPDF Story / python-docx), so the suite
needs no fixture files. See docs/compare_audit_2026-10.md, part E.
"""

import copy
import io

import pymupdf as fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from server.services.processors._diff_engines import paragraph_semantic_diff
from server.services.processors.docx import DocxProcessor, _extract_docx
from server.services.processors.pdf import PDFProcessor, _extract_text_with_pages

FILLER = ('Additional guidance paragraph number {i}: the workstation shall be kept clean and free of foreign '
          'objects, tools shall be counted before and after each operation, and any missing tool shall be '
          'reported immediately to the supervisor in charge.')
NEW_PARAGRAPH = ('NEW PARAGRAPH: this revision introduces a mandatory double check by a second inspector for every '
                 'part classified as flight critical, recorded on the routing card with both signatures.')
TABLE = [['Fastener', 'Min torque', 'Max torque', 'Tool'],
         ['M6 bolt', '10 Nm', '10 Nm', 'Wrench A'],
         ['M8 bolt', '22 Nm', '', 'Wrench B'],
         ['M10 bolt', '', '45 Nm', 'Wrench C']]
RACI = [['Activity', 'Operator', 'Inspector', 'Quality manager'],
        ['Record results', 'X', '', ''],
        ['Release the part', '', 'X', ''],
        ['Approve deviation', '', '', 'X']]


def sections(n_filler=14):
    return [
        ('1. SCOPE', [
            'This procedure defines the inspection requirements applicable to all structural parts manufactured '
            'in the Toulouse plant. It applies to every production order released after the approval date and '
            'covers both metallic and composite components delivered to the final assembly line.',
            'The quality manager shall ensure that the present procedure is known and applied by every operator '
            'involved in the inspection activities described hereafter, including temporary staff.',
        ]),
        ('2. REFERENCES', [
            'The following documents are applicable: EN 9100 revision 2018 for the quality management system, '
            'NAS 410 revision 4 for the qualification of personnel, and the internal instruction QR-1005.',
        ]),
        ('3. INSPECTION', [
            'The operator shall record the inspection results in the quality register before the part is released '
            'to the next production step and shall keep the record available for the customer representative.',
            'The torque applied on each fastener shall be 35 Nm with a tolerance of plus or minus 2 Nm, verified '
            'with a calibrated wrench whose calibration date is checked before each shift by the team leader.',
            'Any part showing a visible defect larger than 0.5 mm shall be isolated in the quarantine area and '
            'identified with a red label until the material review board has decided on its disposition.',
        ]),
        ('3 bis. HOUSEKEEPING', [FILLER.format(i=i) for i in range(1, n_filler + 1)]),
        ('4. RECORDS', [
            'Inspection records shall be archived for a minimum of 10 years in the document management system '
            'and shall remain retrievable within two working days upon request from the authority.',
        ]),
    ]


def edit(secs, title, idx, old, new):
    secs = copy.deepcopy(secs)
    paras = dict(secs)[title]
    assert old in paras[idx]
    paras[idx] = paras[idx].replace(old, new)
    return secs


def make_pdf(secs, footer='DOC QR-2040 - Rev A'):
    """A paginated PDF with a running header and a 'Page i/n' footer."""
    html = ''.join(f'<h2>{t}</h2>' + ''.join(f'<p>{p}</p>' for p in paras) for t, paras in secs)
    story = fitz.Story(html=html, user_css='body{font-family:sans-serif;font-size:11pt} h2{font-size:13pt}')
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    mediabox = fitz.paper_rect('a4')
    more = 1
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(mediabox + (60, 80, -60, -80))
        story.draw(dev)
        writer.end_page()
    writer.close()
    doc = fitz.open('pdf', buf.getvalue())
    for i, page in enumerate(doc, start=1):
        page.insert_text((60, 40), 'LATECOERE - QUALITY PROCEDURE', fontsize=9)
        page.insert_text((60, 810), f'{footer}    Page {i}/{len(doc)}', fontsize=9)
    out = doc.tobytes()
    doc.close()
    return out


def ruled_table_pdf(table, rows_per_page=50):
    """A PDF whose table has drawn borders (the find_tables() path)."""
    doc = fitz.open()
    header, rows = table[0], table[1:]
    for start in range(0, len(rows), rows_per_page):
        page = doc.new_page()
        page.insert_text((60, 70), '5. TABLE', fontsize=13)
        chunk = ([header] if start == 0 else []) + rows[start:start + rows_per_page]
        for r, row in enumerate(chunk):
            for c, val in enumerate(row):
                rect = fitz.Rect(60 + c * 115, 300 + r * 26, 175 + c * 115, 326 + r * 26)
                page.draw_rect(rect, color=(0, 0, 0), width=0.7)
                if val:
                    page.insert_textbox(rect + (4, 6, -4, -2), val, fontsize=9)
    out = doc.tobytes()
    doc.close()
    return out


def make_docx(secs, table=None, mutate=None):
    d = Document()
    for title, paras in secs:
        d.add_heading(title, level=2)
        for p in paras:
            d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    if mutate:
        mutate(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pdf_text(data):
    return _extract_text_with_pages(data)


def docx_text(data):
    return '\n'.join(_extract_docx(data)[0])


def entries(old_text, new_text):
    diff, _ = paragraph_semantic_diff(old_text, new_text, page_label='Page')
    return [line for line in diff.splitlines() if line and not line.startswith('##')]


BASE = sections()


@pytest.fixture(scope='module')
def base_pdf():
    return pdf_text(make_pdf(BASE))


@pytest.fixture(scope='module')
def base_docx():
    return docx_text(make_docx(BASE, TABLE))


# --- PDF ---

def test_pdf_identical_documents_give_no_entry(base_pdf):
    assert entries(base_pdf, pdf_text(make_pdf(BASE))) == []


def test_pdf_ligatures_are_expanded(base_pdf):
    assert 'defines' in base_pdf and 'ﬁ' not in base_pdf


def test_pdf_value_change_is_one_modified(base_pdf):
    out = entries(base_pdf, pdf_text(make_pdf(edit(BASE, '3. INSPECTION', 1, '35 Nm', '40 Nm'))))
    assert len(out) == 1 and '~~35~~ **40**' in out[0]


def test_pdf_negation_is_reported(base_pdf):
    out = entries(base_pdf, pdf_text(make_pdf(edit(BASE, '3. INSPECTION', 0, 'shall record', 'shall not record'))))
    assert len(out) == 1 and '**not**' in out[0]


@pytest.mark.parametrize('n_filler', [14, 60])
def test_pdf_inserted_paragraph_does_not_create_repagination_noise(n_filler):
    """Every page break below the insertion moves; only the insertion is a change."""
    old = sections(n_filler)
    new = copy.deepcopy(old)
    new[0][1].insert(0, NEW_PARAGRAPH)
    out = entries(pdf_text(make_pdf(old)), pdf_text(make_pdf(new)))
    assert len(out) == 1 and out[0].startswith('ADDED')


def test_pdf_paragraph_cut_by_a_page_break_is_one_block():
    text = pdf_text(make_pdf(sections(60)))
    for i in range(1, 61):
        assert sum(f'paragraph number {i}:' in line for line in text.splitlines()) == 1
        line = next(l for l in text.splitlines() if f'paragraph number {i}:' in l)
        assert 'supervisor in charge.' in line


def test_pdf_footer_revision_change_is_collapsed_to_one_entry():
    long = sections(40)  # 3+ pages: the collapse needs _BOILERPLATE_MIN_PAGES repeats
    out = entries(pdf_text(make_pdf(long)), pdf_text(make_pdf(long, footer='DOC QR-2040 - Rev B')))
    assert len(out) == 1 and 'repeated on' in out[0]


def test_pdf_moved_section_is_not_reported_as_changed(base_pdf):
    moved = copy.deepcopy(BASE)
    moved.append(moved.pop(1))
    assert [e for e in entries(base_pdf, pdf_text(make_pdf(moved))) if not e.startswith('RELOCATED')] == []


def test_pdf_ruled_table_rows_are_labelled_with_their_column():
    lines = pdf_text(ruled_table_pdf(TABLE)).splitlines()
    assert 'Fastener: M8 bolt | Min torque: 22 Nm | Tool: Wrench B [Page 1]' in lines
    assert 'Fastener: M10 bolt | Max torque: 45 Nm | Tool: Wrench C [Page 1]' in lines


def test_pdf_mark_moved_to_another_column_is_reported():
    moved = copy.deepcopy(RACI)
    moved[2] = ['Release the part', '', '', 'X']
    out = entries(pdf_text(ruled_table_pdf(RACI)), pdf_text(ruled_table_pdf(moved)))
    assert len(out) == 1
    assert '~~Inspector:~~ **Quality manager:** X' in out[0]


def test_pdf_table_continued_on_next_page_keeps_its_header():
    rows = [['Step %d' % i, 'X' if i % 2 else '', '' if i % 2 else 'X', ''] for i in range(1, 9)]
    text = pdf_text(ruled_table_pdf([RACI[0]] + rows, rows_per_page=4))
    assert 'Activity: Step 7 | Operator: X [Page 2]' in text.splitlines()


def test_pdf_two_column_form_is_not_labelled():
    """A key/value table has no header: labelling would spread one changed value over every row."""
    form = [['Reference', 'QR-2040'], ['Revision', 'A'], ['Owner', 'Quality']]
    assert 'Revision | A [Page 1]' in pdf_text(ruled_table_pdf(form)).splitlines()


def test_pdf_processor_end_to_end_reports_scanned_document():
    def scanned():
        doc = fitz.open()
        doc.new_page()
        out = doc.tobytes()
        doc.close()
        return out

    result = PDFProcessor('structured').build_messages(scanned(), 'a.pdf', scanned(), 'b.pdf')
    assert len(result.warnings) == 2


# --- DOCX ---

def test_docx_identical_documents_give_no_entry(base_docx):
    assert entries(base_docx, docx_text(make_docx(BASE, TABLE))) == []


def test_docx_table_cells_are_labelled_with_their_own_column(base_docx):
    lines = base_docx.splitlines()
    assert any(l.startswith('Fastener: M6 bolt | Min torque: 10 Nm | Max torque: 10 Nm | Tool: Wrench A') for l in lines)
    assert any(l.startswith('Fastener: M10 bolt | Max torque: 45 Nm | Tool: Wrench C') for l in lines)


def test_docx_table_value_change_names_the_right_column(base_docx):
    changed = [row[:] for row in TABLE]
    changed[3][2] = '48 Nm'
    out = entries(base_docx, docx_text(make_docx(BASE, changed)))
    assert len(out) == 1 and 'Max torque: ~~45~~ **48**' in out[0]


def test_docx_mark_moved_to_another_column_is_reported():
    moved = copy.deepcopy(RACI)
    moved[2] = ['Release the part', '', '', 'X']
    out = entries(docx_text(make_docx([], RACI)), docx_text(make_docx([], moved)))
    assert len(out) == 1 and '~~Inspector:~~ **Quality manager:** X' in out[0]


def _append_run(parent, text):
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    parent.append(r)


def test_docx_tracked_insertion_is_read_and_tracked_deletion_is_not(base_docx):
    def mutate(d):
        p = d.paragraphs[1]._p
        ins = OxmlElement('w:ins')
        _append_run(ins, ' The limit is now 50 units.')
        p.append(ins)
        deleted = OxmlElement('w:del')
        r = OxmlElement('w:r')
        dt = OxmlElement('w:delText')
        dt.text = ' Deleted sentence.'
        r.append(dt)
        deleted.append(r)
        p.append(deleted)

    out = entries(base_docx, docx_text(make_docx(BASE, TABLE, mutate)))
    assert len(out) == 1 and '**The limit is now 50 units.**' in out[0]
    assert 'Deleted sentence' not in out[0]


def test_docx_paragraph_inside_a_content_control_is_read(base_docx):
    sentence = 'Content control paragraph: the release certificate shall be signed by the quality manager.'

    def mutate(d):
        p = d.add_paragraph(sentence)
        sdt, content = OxmlElement('w:sdt'), OxmlElement('w:sdtContent')
        d.element.body.remove(p._p)
        content.append(p._p)
        sdt.append(content)
        d.element.body.insert(5, sdt)

    out = entries(base_docx, docx_text(make_docx(BASE, TABLE, mutate)))
    assert out == [next(e for e in out if sentence in e)] and out[0].startswith('ADDED')


def test_docx_text_inside_a_field_is_read(base_docx):
    def mutate(d):
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'REF x')
        _append_run(fld, ' See also instruction QR-2099.')
        d.paragraphs[1]._p.append(fld)

    out = entries(base_docx, docx_text(make_docx(BASE, TABLE, mutate)))
    assert len(out) == 1 and 'QR-2099' in out[0]


def test_docx_nested_table_is_read(base_docx):
    def mutate(d):
        nested = d.tables[0].cell(1, 3).add_table(rows=1, cols=2)
        nested.cell(0, 0).text = 'Calibration'
        nested.cell(0, 1).text = 'every 6 months'

    out = entries(base_docx, docx_text(make_docx(BASE, TABLE, mutate)))
    assert len(out) == 1 and 'every 6 months' in out[0]


def test_docx_table_of_contents_entries_are_skipped():
    def mutate(d):
        style = d.styles.add_style('TOC 1', 1)
        d.add_paragraph('3. INSPECTION\t4', style=style)

    assert '3. INSPECTION 4' not in docx_text(make_docx(BASE, TABLE, mutate))


def test_docx_processor_end_to_end_builds_the_llm_message():
    old = make_docx(BASE, TABLE)
    new = make_docx(edit(BASE, '4. RECORDS', 0, '10 years', '15 years'), TABLE)
    result = DocxProcessor('structured').build_messages(old, 'a.docx', new, 'b.docx')
    text = result.messages[-1]['content'][0]['text']
    assert '~~10~~ **15** years' in text
    assert text.count('MODIFIED [') == 1 and 'ADDED [' not in text and 'REMOVED [' not in text


# --- Both formats ---

@pytest.mark.parametrize('make, read', [(make_pdf, pdf_text), (lambda s: make_docx(s, TABLE), docx_text)])
def test_paragraph_split_or_merged_is_not_a_change(make, read):
    split = copy.deepcopy(BASE)
    para = split[0][1][0]
    cut = para.index(' It applies')
    split[0][1][0:1] = [para[:cut], para[cut + 1:]]
    whole_text, split_text = read(make(BASE)), read(make(split))
    assert entries(whole_text, split_text) == []
    assert entries(split_text, whole_text) == []


@pytest.mark.parametrize('make, read', [(make_pdf, pdf_text), (lambda s: make_docx(s, TABLE), docx_text)])
def test_renumbered_headings_are_paired_not_removed_and_added(make, read):
    new = copy.deepcopy(BASE)
    new.insert(2, ('3. TRAINING', ['Every inspector shall complete the internal training course before performing '
                                   'any inspection alone, and the record shall be kept by human resources.']))
    new[3] = ('4. INSPECTION', new[3][1])
    new[5] = ('5. RECORDS', new[5][1])
    out = entries(read(make(BASE)), read(make(new)))
    assert not any(e.startswith('REMOVED') for e in out)
    assert any('~~3.~~ **4.** INSPECTION' in e for e in out)
    assert any('~~4.~~ **5.** RECORDS' in e for e in out)


@pytest.mark.parametrize('make, read', [(make_pdf, pdf_text), (lambda s: make_docx(s, TABLE), docx_text)])
def test_several_edits_including_a_moved_and_modified_paragraph(make, read):
    new = edit(BASE, '3. INSPECTION', 1, '35 Nm', '40 Nm')
    new = edit(new, '4. RECORDS', 0, '10 years', '15 years')
    new = edit(new, '2. REFERENCES', 0, 'NAS 410 revision 4', 'NAS 410 revision 5')
    moved = new[2][1].pop(2)
    new[4][1].append(moved.replace('0.5 mm', '0.3 mm'))
    out = entries(read(make(BASE)), read(make(new)))
    assert len(out) == 4 and all(e.startswith('MODIFIED') for e in out)
    for marker in ('~~35~~ **40**', '~~10~~ **15**', '~~4~~ **5**', '~~0.5~~ **0.3**'):
        assert any(marker in e for e in out)


def test_value_that_looks_like_a_heading_number_stays_visible():
    out = entries('Gap table\n2.5 mm maximum gap [Page 1]', 'Gap table\n3.5 mm maximum gap [Page 1]')
    assert out == ['MODIFIED [Page 1]: ~~2.5~~ **3.5** mm maximum gap']
