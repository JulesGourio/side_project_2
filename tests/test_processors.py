"""PDF, DOCX, PPTX and XML processors: LLM messages, image pairs, warnings and metadata of small generated files."""

import io

import pymupdf
import pytest
from docx import Document
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from server.services.compare.processors.docx import DocxProcessor
from server.services.compare.processors.pdf import PDFProcessor
from server.services.compare.processors.pptx import PptxProcessor
from server.services.compare.processors.xml import XMLProcessor, _extract_xml

SCOPE = '1. SCOPE'
OPERATOR = 'The operator shall record every torque value measured on the primary fastener assembly before delivery.'
SUPPLIER = 'The supplier must deliver the part within ten working days after approval of the drawing by the customer.'
INSPECTION = '2. INSPECTION'
FREQUENCY = 'Inspection frequency is one part in ten for the first batch and one part in fifty afterwards.'
RECORDS = 'Records are kept for five years in the quality register of the plant concerned.'
SIGNATURE = 'A new paragraph states that the inspector signs each page of the traveller before release.'

OLD = [SCOPE, OPERATOR, SUPPLIER, INSPECTION, FREQUENCY, RECORDS]
NEW = [SCOPE, OPERATOR.replace('delivery', 'shipment'), SUPPLIER.replace('must', 'may'), INSPECTION,
       'Inspection frequency is one part in ten for every batch.', RECORDS, SIGNATURE]


def _png(edited: bool = False) -> bytes:
    """A small picture; the edited one carries a diagonal line.

    :param edited: draws the diagonal line
    :return: PNG bytes
    """
    img = Image.new('RGB', (120, 90), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.rectangle([10, 10, 60, 40], fill=(40, 40, 40))
    draw.rectangle([70, 50, 110, 80], fill=(160, 160, 160))
    if edited:
        draw.line([0, 0, 119, 89], fill=(0, 0, 0), width=4)
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    return buf.getvalue()


def _pdf(lines: list[str], edited_image: bool | None = False) -> bytes:
    """A one-page PDF with one text line per row and, unless ``edited_image`` is None, one picture.

    :param lines: text lines, written top to bottom
    :param edited_image: whether the picture is the edited one; None for no picture
    :return: PDF bytes
    """
    doc = pymupdf.open()
    page = doc.new_page()
    for i, line in enumerate(lines):
        page.insert_text((72, 72 + 18 * i), line[:95], fontsize=10)
    if edited_image is not None:
        page.insert_image(pymupdf.Rect(72, 300, 192, 390), stream=_png(edited_image))
    return doc.tobytes()


def _docx(lines: list[str], delay: str = '3 days', edited_image: bool | None = False) -> bytes:
    """A DOCX with one paragraph per line, a 2x3 table whose last cell is ``delay`` and an optional picture.

    :param lines: paragraphs
    :param delay: text of the last table cell
    :param edited_image: whether the picture is the edited one; None for no picture
    :return: DOCX bytes
    """
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    table = doc.add_table(rows=2, cols=3)
    for col, text in enumerate(('Activity', 'Owner', 'Delay')):
        table.cell(0, col).text = text
    for col, text in enumerate(('Release', 'Quality', delay)):
        table.cell(1, col).text = text
    if edited_image is not None:
        doc.add_picture(io.BytesIO(_png(edited_image)))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx(slides: list[list[str]], edited_image: bool | None = False) -> bytes:
    """A PPTX with one slide per list (title first, body lines after) and a picture on the last slide.

    :param slides: title and body lines of each slide
    :param edited_image: whether the picture is the edited one; None for no picture
    :return: PPTX bytes
    """
    prs = Presentation()
    for chunk in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = chunk[0]
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(3))
        box.text_frame.text = '\n'.join(chunk[1:])
    if edited_image is not None:
        slide.shapes.add_picture(io.BytesIO(_png(edited_image)), Inches(1), Inches(5), Inches(2), Inches(1.5))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _xml(lines: list[str]) -> bytes:
    """An XML file with one ``<item><name>`` per line.

    :param lines: element texts
    :return: XML bytes
    """
    body = ''.join(f'<item id="{i}"><name>{line}</name></item>' for i, line in enumerate(lines))
    return f'<?xml version="1.0"?><root>{body}</root>'.encode()


def _user_text(result) -> str:
    """Text block of the user message.

    :param result: ``ProcessResult``
    :return: the first text block of the last message
    """
    return result.messages[-1]['content'][0]['text']


def _changes(text: str) -> list[str]:
    """Kind (MODIFIED, ADDED, REMOVED) of each change line of a diff.

    :param text: text block of the user message
    :return: the kinds, in order
    """
    return [line.split(' ')[0] for line in text.splitlines() if line.split(' ')[0] in ('MODIFIED', 'ADDED', 'REMOVED')]


def _image_blocks(result) -> list[dict]:
    """Image blocks attached to the user message.

    :param result: ``ProcessResult``
    :return: the ``image_url`` blocks
    """
    return [b for b in result.messages[-1]['content'] if b['type'] == 'image_url']


# --- XML ---

def test_xml_changed_value_is_reported_with_its_item_number():
    new = list(OLD)
    new[4] = FREQUENCY.replace('fifty', 'forty')
    result = XMLProcessor().build_messages(_xml(OLD), 'a.xml', _xml(new), 'b.xml')
    text = _user_text(result)
    assert '## 2. INSPECTION' in text
    assert ('MODIFIED [Item 5]: Inspection frequency is one part in ten for the first batch and one part in '
            '~~fifty~~ **forty** afterwards.') in text
    assert _changes(text) == ['MODIFIED']
    assert 'TEXT CHANGES' in text and 'VISUAL CHANGES' not in text


def test_xml_wording_swaps_and_whitespace_are_not_reported():
    noisy = list(OLD)
    noisy[2] = '  ' + SUPPLIER.replace('must', 'shall').replace(' ', '  ') + ' '
    result = XMLProcessor().build_messages(_xml(OLD), 'a.xml', _xml(noisy), 'b.xml')
    assert _user_text(result).startswith('No significant text changes detected.')
    assert _changes(_user_text(result)) == []


def test_xml_added_and_modified_items_across_sections():
    result = XMLProcessor().build_messages(_xml(OLD), 'a.xml', _xml(NEW), 'b.xml')
    text = _user_text(result)
    assert text.index('## 1. SCOPE') < text.index('## 2. INSPECTION')
    assert 'before ~~delivery.~~ **shipment.**' in text
    assert 'The supplier ~~must~~ **may** deliver' in text
    assert f'ADDED [Item 7]: {SIGNATURE}' in text
    assert _changes(text) == ['MODIFIED', 'MODIFIED', 'MODIFIED', 'ADDED']


def test_xml_text_and_tail_nodes_are_numbered_in_document_order():
    lines = _extract_xml(b'<r>  <a>one</a>tail <b> <c>two</c></b>  </r>')
    assert lines == ['one [Item 1]', 'tail [Item 2]', 'two [Item 3]']


def test_xml_malformed_file_raises_instead_of_reporting_no_change():
    with pytest.raises(ValueError, match='not well formed'):
        XMLProcessor().build_messages(b'<a><b></a>', 'a.xml', _xml(OLD), 'b.xml')


@pytest.mark.parametrize('method, system_marker', [
    ('standard', 'document comparison analyst'),
    ('structured', 'Output ONLY valid JSON'),
])
def test_xml_standard_and_structured_use_their_own_system_prompt(method, system_marker):
    result = XMLProcessor(method).build_messages(_xml(OLD), 'a.xml', _xml(NEW), 'b.xml', 'ignored custom prompt')
    assert [m['role'] for m in result.messages] == ['system', 'user']
    assert system_marker in result.messages[0]['content']
    assert 'ignored custom prompt' not in result.messages[0]['content']
    assert (result.metadata.file_type, result.metadata.method) == ('xml', method)
    assert (result.metadata.old_name, result.metadata.new_name) == ('a.xml', 'b.xml')
    assert result.image_pairs == [] and result.warnings == []


def test_xml_comparative_groups_changes_by_section_with_the_custom_prompt():
    result = XMLProcessor('comparative').build_messages(_xml(OLD), 'a.xml', _xml(NEW), 'b.xml', 'custom system prompt')
    assert result.messages[0] == {'role': 'system', 'content': 'custom system prompt'}
    text = _user_text(result)
    assert 'CHANGES BY SECTION' in text and '## 2. INSPECTION' in text
    assert 'REMOVED: ' + FREQUENCY + ' [Item 5]' in text
    assert 'ADDED: Inspection frequency is one part in ten for every batch. [Item 5]' in text
    assert f'ADDED: {SIGNATURE} [Item 7]' in text
    assert result.metadata.method == 'comparative'


def test_comparative_without_prompt_sends_only_the_user_message():
    result = XMLProcessor('comparative').build_messages(_xml(OLD), 'a.xml', _xml(NEW), 'b.xml')
    assert [m['role'] for m in result.messages] == ['user']


# --- PDF ---

def test_pdf_text_changes_are_tagged_with_their_page():
    result = PDFProcessor().build_messages(_pdf(OLD), 'a.pdf', _pdf(NEW), 'b.pdf')
    text = _user_text(result)
    assert text.index('## 1. SCOPE') < text.index('## 2. INSPECTION')
    assert 'The supplier ~~must~~ **may** deliver' in text
    assert ('MODIFIED [Page 1]: Inspection frequency is one part in ten for ~~the first batch and one part in '
            'fifty afterwards.~~ **every batch.**') in text
    assert f'ADDED [Page 1]: {SIGNATURE}' in text
    assert (result.metadata.file_type, result.metadata.method) == ('pdf', 'standard')
    assert result.warnings == []


def test_pdf_changed_picture_is_attached_as_an_old_new_pair():
    result = PDFProcessor().build_messages(_pdf(OLD, False), 'a.pdf', _pdf(OLD, True), 'b.pdf')
    assert _user_text(result).startswith('No significant text changes detected.')
    assert len(result.image_pairs) == 1
    pair = result.image_pairs[0]
    assert (pair['status'], pair['old_page'], pair['new_page']) == ('modified', 1, 1)
    assert pair['old_b64'] and pair['new_b64'] and pair['old_b64'] != pair['new_b64']
    assert len(_image_blocks(result)) == 2
    assert '1 image(s) MODIFIED (OLD → NEW)' in str(result.messages[-1]['content'])


def test_pdf_identical_picture_is_not_attached():
    result = PDFProcessor().build_messages(_pdf(OLD), 'a.pdf', _pdf(OLD), 'b.pdf')
    assert result.image_pairs == [] and _image_blocks(result) == []


def test_pdf_without_text_layer_warns_for_both_files():
    blank = pymupdf.open()
    blank.new_page()
    data = blank.tobytes()
    result = PDFProcessor().build_messages(data, 'a.pdf', data, 'b.pdf')
    assert len(result.warnings) == 2
    assert result.warnings[0].startswith('OLD document yielded almost no extractable text')
    assert result.warnings[1].startswith('NEW document yielded almost no extractable text')


# --- DOCX ---

def test_docx_table_row_is_labelled_with_its_header_and_the_changed_number_is_reported():
    old = _docx(OLD, '3 days', None)
    new = _docx(OLD, '5 days', None)
    result = DocxProcessor().build_messages(old, 'a.docx', new, 'b.docx')
    text = _user_text(result)
    assert 'MODIFIED [Page 1]: Activity: Release | Owner: Quality | Delay: ~~3~~ **5** days' in text
    assert _changes(text) == ['MODIFIED']
    assert (result.metadata.file_type, result.metadata.method) == ('docx', 'standard')


def test_docx_paragraph_changes_and_changed_picture():
    result = DocxProcessor().build_messages(_docx(OLD), 'a.docx', _docx(NEW, edited_image=True), 'b.docx')
    text = _user_text(result)
    assert 'before ~~delivery.~~ **shipment.**' in text
    assert f'ADDED [Page 1]: {SIGNATURE}' in text
    assert [p['status'] for p in result.image_pairs] == ['modified']
    assert len(_image_blocks(result)) == 2


def test_docx_comparative_tags_page_and_paragraph():
    result = DocxProcessor('comparative').build_messages(_docx(OLD, edited_image=None), 'a.docx',
                                                         _docx(NEW, edited_image=None), 'b.docx')
    text = _user_text(result)
    assert f'REMOVED: {FREQUENCY} [Page 1, Para 5]' in text
    assert 'ADDED: Inspection frequency is one part in ten for every batch. [Page 1, Para 5]' in text
    assert result.image_pairs == []


# --- PPTX ---

def test_pptx_changes_are_tagged_with_their_slide_and_picture_with_its_page():
    old = _pptx([OLD[:3], OLD[3:]], False)
    new = _pptx([NEW[:3], NEW[3:]], True)
    result = PptxProcessor().build_messages(old, 'a.pptx', new, 'b.pptx')
    text = _user_text(result)
    assert 'MODIFIED [Slide 1]: The supplier ~~must~~ **may** deliver' in text
    assert 'MODIFIED [Slide 2]: Inspection frequency is one part in ten for ~~the first batch' in text
    assert f'ADDED [Slide 2]: {SIGNATURE}' in text
    assert [(p['status'], p['old_page'], p['new_page']) for p in result.image_pairs] == [('modified', 2, 2)]
    assert len(_image_blocks(result)) == 2
    assert (result.metadata.file_type, result.metadata.method) == ('pptx', 'standard')


def test_pptx_comparative_groups_by_section_and_keeps_the_slide():
    old = _pptx([OLD[:3], OLD[3:]], None)
    new = _pptx([NEW[:3], NEW[3:]], None)
    result = PptxProcessor('comparative').build_messages(old, 'a.pptx', new, 'b.pptx', 'custom system prompt')
    text = _user_text(result)
    assert 'CHANGES BY SECTION' in text and '## 2. INSPECTION' in text
    assert f'ADDED: {SIGNATURE} [Slide 2]' in text
    assert result.messages[0]['content'] == 'custom system prompt'
