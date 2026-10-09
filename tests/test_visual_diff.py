"""Tests for the visual-diff pipeline: two-tier image identity, extraction
warnings, and DOCX estimated page numbers for images.

Coverage:
  - _compute_image_diff_pairs : fine-hash (16x16) second-tier identity — a
    subtly edited diagram becomes a MODIFIED pair instead of "unchanged";
    recompression noise and legacy entries (no hash16) stay "unchanged".
  - extraction_warnings       : scanned/empty documents flagged.
  - _extract_docx             : image positions remapped to estimated pages.
"""

import base64
import io

from PIL import Image as PILImage

from server.services.processors._diff_engines import (
    _compute_image_diff_pairs,
    extraction_warnings,
)
from server.services.processors.docx import _extract_docx


# --- Helpers ---

def _jpeg_b64(color=(255, 255, 255), size=(64, 64)) -> str:
    img = PILImage.new('RGB', size, color)
    buf = io.BytesIO()
    img.save(buf, format='JPEG', quality=85)
    return base64.b64encode(buf.getvalue()).decode('ascii')


_DHASH_KEY = '0' * 64          # same coarse hash on both sides
_FINE_BASE = '0' * 256


def _fine_variant(bits_flipped: int) -> str:
    return '1' * bits_flipped + '0' * (256 - bits_flipped)


def _img(fine_hash=None, page=1, b64=None):
    meta = {'b64': b64 or _jpeg_b64(), 'pages': [page]}
    if fine_hash is not None:
        meta['hash16'] = fine_hash
    return meta


# --- Two-tier image identity ---

def test_identical_images_are_not_reported():
    old = {_DHASH_KEY: _img(fine_hash=_FINE_BASE)}
    new = {_DHASH_KEY: _img(fine_hash=_FINE_BASE)}
    modified, orientation, removed, added = _compute_image_diff_pairs(old, new, show_position=True)
    assert (modified, orientation, removed, added) == ([], [], [], [])


def test_recompression_noise_stays_unchanged():
    # Small fine-hash distance (8 <= threshold 12) = same content re-saved
    old = {_DHASH_KEY: _img(fine_hash=_FINE_BASE)}
    new = {_DHASH_KEY: _img(fine_hash=_fine_variant(8))}
    modified, orientation, removed, added = _compute_image_diff_pairs(old, new, show_position=True)
    assert (modified, orientation, removed, added) == ([], [], [], [])


def test_subtle_diagram_edit_becomes_modified_pair():
    # Coarse hashes identical (would have been missed before) but fine hashes
    # differ well past the threshold — must surface as a MODIFIED pair.
    b64 = _jpeg_b64()
    old = {_DHASH_KEY: _img(fine_hash=_FINE_BASE, b64=b64)}
    new = {_DHASH_KEY: _img(fine_hash=_fine_variant(20), b64=b64)}
    modified, orientation, removed, added = _compute_image_diff_pairs(old, new, show_position=True)
    assert modified == [(_DHASH_KEY, _DHASH_KEY)]
    assert removed == [] and added == []


def test_legacy_entries_without_fine_hash_keep_old_behaviour():
    old = {_DHASH_KEY: _img()}
    new = {_DHASH_KEY: _img()}
    modified, orientation, removed, added = _compute_image_diff_pairs(old, new, show_position=True)
    assert (modified, orientation, removed, added) == ([], [], [], [])


def test_single_leftover_pair_on_same_page_is_a_replacement():
    # Visually unrelated hashes (hamming 32 > gates) but one removed + one
    # added on the same page → figure replaced in place → MODIFIED pair.
    other_key = '1' * 32 + '0' * 32
    old = {_DHASH_KEY: _img(page=6)}
    new = {other_key: _img(page=6)}
    modified, orientation, removed, added = _compute_image_diff_pairs(old, new, show_position=True)
    assert modified == [(_DHASH_KEY, other_key)]
    assert removed == [] and added == []


def test_multiple_leftovers_on_same_page_are_not_force_paired():
    key_b = '1' * 32 + '0' * 32
    key_c = '0' * 32 + '1' * 32
    old = {_DHASH_KEY: _img(page=6), key_b: _img(page=6)}
    new = {key_c: _img(page=6)}
    modified, orientation, removed, added = _compute_image_diff_pairs(old, new, show_position=True)
    assert modified == []
    assert len(removed) == 2 and len(added) == 1


# --- Extraction warnings ---

def test_warnings_for_scanned_documents():
    long_text = ('The operator shall verify the torque value at every step. ' * 20)
    assert extraction_warnings(long_text, long_text) == []

    both = extraction_warnings('', '')
    assert len(both) == 2

    one = extraction_warnings(long_text, 'Page 3')
    assert len(one) == 1
    assert one[0].startswith('NEW')


# --- DOCX — image positions remapped to estimated pages ---

def _png_stream(pattern: str) -> io.BytesIO:
    """Structured test images — solid colors all hash to zero gradients, so
    use patterns that produce distinct dhashes."""
    img = PILImage.new('RGB', (48, 48), 'white')
    px = img.load()
    for x in range(48):
        for y in range(48):
            if pattern == 'left-half' and x < 24:
                px[x, y] = (0, 0, 0)
            elif pattern == 'checker' and (x // 6 + y // 6) % 2 == 0:
                px[x, y] = (0, 0, 0)
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return buf


def test_docx_images_get_estimated_pages():
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Inches

    doc = Document()
    doc.add_paragraph('Introduction paragraph with some meaningful text content.')
    doc.add_picture(_png_stream('left-half'), width=Inches(1))
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run('Section two starts here with more meaningful text content.')
    doc.add_picture(_png_stream('checker'), width=Inches(1))

    buf = io.BytesIO()
    doc.save(buf)
    texts, images = _extract_docx(buf.getvalue())

    assert len(images) == 2
    pages = sorted(meta['pages'][0] for meta in images.values())
    assert pages == [1, 2]
    assert all('hash16' in meta for meta in images.values())
    assert any('[Page 2' in t for t in texts)  # page break was honoured


# --- DOCX processor end-to-end — operator-style document (text + image change) ---

def _operator_docx(torque: str, image_pattern: str) -> bytes:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_paragraph('1. ASSEMBLY INSTRUCTIONS')
    doc.add_paragraph(f'The operator shall apply a torque of {torque} N·m on connector J3 before inspection.')
    doc.add_picture(_png_stream(image_pattern), width=Inches(1))
    doc.add_paragraph('2. FINAL CHECK')
    doc.add_paragraph('Record the result on the routing card after each assembly operation.')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_processor_reports_text_and_image_changes():
    from server.services.processors.docx import DocxProcessor

    old_bytes = _operator_docx('35', 'left-half')
    new_bytes = _operator_docx('40', 'checker')

    result = DocxProcessor('structured').build_messages(old_bytes, 'old.docx', new_bytes, 'new.docx')

    text_block = result.messages[1]['content'][0]['text']
    assert 'MODIFIED' in text_block
    assert '~~35~~' in text_block and '**40**' in text_block

    # The changed picture must surface as image blocks for the LLM…
    assert any(b.get('type') == 'image_url' for b in result.messages[1]['content'])
    # …and as UI pairs carrying (estimated) page positions — new for DOCX.
    assert result.image_pairs
    assert all((p['old_page'] or p['new_page']) == 1 for p in result.image_pairs)
    assert result.warnings == []  # normal text density — no scanned-doc warning
