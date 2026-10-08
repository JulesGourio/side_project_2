"""Tests for impact-search change extraction, query derivation and per-document aggregation.

Coverage:
  - changes_to_queries : numbered changes, one query per searched change, overflow packing
  - _aggregate_docs    : change-id tracking, ranking, self-exclusion, language variants
  - _judged_doc        : passage mapping, quote highlighting, status derivation
"""

import json

from server.services.impact_queries import changes_to_queries
from server.services.vector_search import (
    _aggregate_docs, _exclusion_keys, _find_quote, _judged_doc, _split_prefix, sort_documents,
)


STRUCTURED = json.dumps([
    {'section': '4.3 Torque', 'type': 'Value changed', 'criticality': 'Medium',
     'before': '35 N·m', 'after': '40 N·m', 'rationale': 'J3 connector torque limit raised.'},
    {'section': '2.1 References', 'type': 'Reference updated', 'criticality': 'High',
     'before': 'NAS 410 Rev 2', 'after': 'NAS 410 Rev 3', 'rationale': ''},
    {'section': 'TOC', 'type': 'Editorial/Structural', 'criticality': 'Low',
     'before': 'section 5', 'after': 'section 6', 'rationale': ''},
    {'section': '6.2 Inspection', 'type': 'Requirement added', 'criticality': 'Low',
     'before': '--', 'after': 'visual inspection before assembly', 'rationale': ''},
])

MARKDOWN = """\
## 4.3 Torque

- Torque limit changed from **35 N·m** to **40 N·m**

## 2.1 References

- NAS 410 revision updated to Rev 3
"""


# --- changes_to_queries ---

def test_structured_changes_keep_table_order_and_ids():
    out = changes_to_queries(STRUCTURED)
    assert out['source'] == 'structured'
    assert [c['id'] for c in out['changes']] == ['C1', 'C2', 'C3', 'C4']
    assert out['changes'][0]['before'] == '35 N·m' and out['changes'][0]['after'] == '40 N·m'


def test_editorial_rows_kept_but_not_searched():
    out = changes_to_queries(STRUCTURED)
    assert out['changes'][2]['searched'] is False
    searched_ids = [cid for q in out['queries'] for cid in q['change_ids']]
    assert 'C3' not in searched_ids
    assert len(out['queries']) == 3


def test_one_query_per_change_high_criticality_first():
    out = changes_to_queries(STRUCTURED)
    assert out['queries'][0]['change_ids'] == ['C2']
    assert all(len(q['change_ids']) == 1 for q in out['queries'])
    assert not any('{' in q['text'] for q in out['queries'])


def test_added_rendering():
    out = changes_to_queries(STRUCTURED)
    added = [q['text'] for q in out['queries'] if 'added:' in q['text']]
    assert added and 'visual inspection' in added[0]


def test_editorial_only_is_still_searched():
    text = json.dumps([{'section': 'TOC', 'type': 'Editorial/Structural',
                        'criticality': 'Low', 'before': 'a', 'after': 'b', 'rationale': ''}])
    out = changes_to_queries(text)
    assert len(out['queries']) == 1


def test_empty_array_means_no_changes():
    out = changes_to_queries('[]')
    assert out['queries'] == [] and out['changes'] == []


def test_markdown_one_change_per_section():
    out = changes_to_queries(MARKDOWN)
    assert out['source'] == 'markdown'
    assert [c['section'] for c in out['changes']] == ['4.3 Torque', '2.1 References']
    assert out['queries'][0]['text'].startswith('4.3 Torque:')
    assert '**' not in out['queries'][0]['text']


def test_markdown_no_changes_detected():
    assert changes_to_queries('**No significant changes detected.**')['queries'] == []


def test_raw_blob_single_change():
    out = changes_to_queries('The torque changed from 35 to 40 N·m.')
    assert out['source'] == 'raw'
    assert out['queries'][0]['change_ids'] == ['C1']


def test_overflow_packs_remaining_changes_without_dropping_any():
    rows = [{'section': f'S{i}', 'type': 'Value changed', 'criticality': 'Medium',
             'before': f'{i}', 'after': f'{i + 1}', 'rationale': ''} for i in range(40)]
    out = changes_to_queries(json.dumps(rows), max_queries=30)
    assert len(out['queries']) <= 30
    ids = [cid for q in out['queries'] for cid in q['change_ids']]
    assert sorted(ids) == sorted(f'C{i}' for i in range(1, 41))


def test_overflow_groups_changes_of_the_same_section():
    rows = [{'section': f'S{i % 10}', 'type': 'Value changed', 'criticality': 'Medium',
             'before': f'{i}', 'after': f'{i + 1}', 'rationale': ''} for i in range(40)]
    out = changes_to_queries(json.dumps(rows), max_queries=30)
    assert len(out['queries']) == 10
    assert out['queries'][0]['change_ids'] == ['C1', 'C11', 'C21', 'C31']


def test_overflow_splits_a_section_too_long_for_one_query():
    rows = [{'section': 'S', 'type': 'Value changed', 'criticality': 'Medium',
             'before': 'x' * 300, 'after': f'{i}', 'rationale': ''} for i in range(40)]
    out = changes_to_queries(json.dumps(rows), max_queries=30)
    assert 1 < len(out['queries']) <= 30
    assert all(len(q['text']) <= 2000 for q in out['queries'])
    assert sum(len(q['change_ids']) for q in out['queries']) == 40


def test_empty_input():
    assert changes_to_queries('')['queries'] == []


# --- _aggregate_docs ---

def _chunk(iddoc, ref, score, change_ids=('C1',), text='chunk text', chunk_id='1-000001'):
    return {'IDDOC': iddoc, 'REF': ref, 'division': 'AS', 'url': '', 'semantic_headers': '',
            'chunk_text': text, 'chunk_id': chunk_id, 'score': score, '_change_ids': list(change_ids)}


def test_same_chunk_from_several_changes_is_kept_once():
    docs = _aggregate_docs([
        _chunk('D1', 'REF-001', 0.9, ['C1']),
        _chunk('D1', 'REF-001', 0.8, ['C2']),
    ])
    assert len(docs[0]['passages']) == 1
    assert docs[0]['passages'][0]['change_ids'] == ['C1', 'C2']
    assert docs[0]['change_ids'] == ['C1', 'C2']


def test_docs_matched_by_more_changes_rank_first():
    docs = _aggregate_docs([
        _chunk('D1', 'REF-001', 0.6, ['C1'], chunk_id='1-000001'),
        _chunk('D1', 'REF-001', 0.6, ['C2'], chunk_id='1-000002'),
        _chunk('D2', 'REF-002', 0.99, ['C1'], chunk_id='2-000001'),
    ])
    assert docs[0]['ref'] == 'REF-001'


def test_passages_in_document_order():
    docs = _aggregate_docs([
        _chunk('D1', 'REF-001', 0.9, chunk_id='1-000007', text='late'),
        _chunk('D1', 'REF-001', 0.5, chunk_id='1-IMG-002', text='image'),
        _chunk('D1', 'REF-001', 0.8, chunk_id='1-000002', text='early'),
    ])
    assert [p['text'] for p in docs[0]['passages']] == ['early', 'late', 'image']


def test_language_variants_merged():
    docs = _aggregate_docs([
        _chunk('D1', 'GO-1316_FR', 0.9, chunk_id='1-000001'),
        _chunk('D2', 'GO-1316_GB', 0.8, ['C2'], chunk_id='2-000001'),
    ])
    assert len(docs) == 1
    assert docs[0]['variants'][0]['ref'] == 'GO-1316_GB'
    assert docs[0]['change_ids'] == ['C1', 'C2']


def test_compared_document_is_excluded():
    excluded = _exclusion_keys(['PR-2207_FR rev G.docx', ''], ['PR-2207_FR', 'MI-0815'])
    docs = _aggregate_docs([
        _chunk('D1', 'PR-2207_FR', 0.99, chunk_id='1-000001'),
        _chunk('D2', 'MI-0815', 0.5, chunk_id='2-000001'),
    ], excluded)
    assert [d['ref'] for d in docs] == ['MI-0815']


def test_prefix_stripped_and_metadata_read():
    body, meta = _split_prefix(
        '[Source: MI-0611 | Title: Porte A320 | Division: AS | Category: Prod | Date de diffusion: 2016-03-01]\n\nTexte'
    )
    assert body == 'Texte'
    assert meta == {'title': 'Porte A320', 'doc_date': '2016-03-01'}


# --- _judged_doc ---

def _candidate():
    return _aggregate_docs([
        _chunk('D1', 'MI-0815', 0.9, chunk_id='1-000001',
               text='[Source: MI-0815 | Title: Cadre | Division: AS | Category: X | Date de diffusion: 2016-01-01]\n\n'
                    '[4.2 Fixation] Serrer au couple de\n12 N·m.'),
        _chunk('D1', 'MI-0815', 0.8, chunk_id='1-000002', text='Autre passage.'),
    ])[0]


def test_judged_doc_maps_passages_and_highlights_quote():
    doc = _judged_doc(_candidate(), {
        'impacted': True, 'confidence': 'high', 'reason': 'r',
        'passages': [{'passage': 1, 'changes': ['C1'], 'section': '4.2 Fixation',
                      'quote': 'couple de 12 N·m', 'explanation': 'e'},
                     {'passage': 9, 'changes': ['C1']}],
    }, '2018-01-01')
    assert doc['status'] == 'impacted' and doc['archive'] is True
    assert len(doc['passages']) == 1
    p = doc['passages'][0]
    start, end = p['highlight']
    assert p['text'][start:end] == 'couple de\n12 N·m'
    assert doc['sections'] == ['4.2 Fixation']


def test_low_confidence_means_check_either_way():
    assert _judged_doc(_candidate(), {'impacted': False, 'confidence': 'low'}, '')['status'] == 'check'


def test_not_impacted_drops_passages():
    doc = _judged_doc(_candidate(), {'impacted': False, 'confidence': 'high',
                                     'passages': [{'passage': 1, 'changes': ['C1']}]}, '')
    assert doc['status'] == 'not_impacted' and doc['passages'] == []


def test_quote_not_found_gives_no_highlight():
    assert _find_quote('abc def', 'xyz') is None


def test_sort_documents_by_status():
    docs = [{'status': 'not_impacted'}, {'status': 'error'}, {'status': 'impacted'}, {'status': 'check'}]
    assert [d['status'] for d in sort_documents(docs)] == ['impacted', 'check', 'not_impacted', 'error']


def test_impact_excel_has_one_row_per_passage():
    import io
    import openpyxl
    from server.services.export_helpers import _build_impact_excel_bytes

    doc = _judged_doc(_candidate(), {
        'impacted': True, 'confidence': 'high', 'reason': 'r',
        'passages': [{'passage': 1, 'changes': ['C1'], 'quote': '12 N·m'},
                     {'passage': 2, 'changes': ['C1']}],
    }, '2018-01-01')
    wb = openpyxl.load_workbook(io.BytesIO(_build_impact_excel_bytes({
        'changes': changes_to_queries(STRUCTURED)['changes'], 'documents': [doc],
        'not_judged': [{'ref': 'X-1', 'judged': False}],
    })))
    assert wb.sheetnames == ['Passages', 'Documents', 'Changes']
    assert wb['Passages'].max_row == 3
    assert wb['Documents'].max_row == 2
    assert wb['Passages']['H2'].value.startswith('C1: 4.3 Torque')
