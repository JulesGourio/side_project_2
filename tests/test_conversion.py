"""Document conversion: .doc text (antiword, olefile), .docx HTML, presentation slide images."""

import io
import shutil
import subprocess
from types import SimpleNamespace

import pymupdf
import pytest
from docx import Document
from pptx import Presentation

from server.services.compare import conversion
from server.services.compare.soffice import SofficeUnavailable


def _pdf(pages: int) -> bytes:
    doc = pymupdf.open()
    for number in range(pages):
        doc.new_page().insert_text((72, 72), f'Slide {number + 1}')
    return doc.tobytes()


def _is_png(data: bytes) -> bool:
    return data.startswith(b'\x89PNG')


def test_every_page_of_the_pdf_becomes_one_image_in_slide_order(monkeypatch):
    monkeypatch.setattr(conversion, 'convert_docx_to_pdf', lambda content, suffix: _pdf(3))
    images = conversion.pptx_to_slide_images(b'deck', 'deck.pptx')
    assert len(images) == 3 and all(_is_png(image) for image in images)
    assert len({image for image in images}) == 3


def test_the_file_extension_is_given_to_libreoffice(monkeypatch):
    seen = {}
    monkeypatch.setattr(conversion, 'convert_docx_to_pdf', lambda content, suffix: seen.update(suffix=suffix) or _pdf(1))
    conversion.pptx_to_slide_images(b'deck', 'old.ppt')
    assert seen['suffix'] == '.ppt'


def test_slides_beyond_the_cap_are_not_rendered(monkeypatch):
    monkeypatch.setattr(conversion, '_MAX_SLIDE_IMAGES', 2)
    monkeypatch.setattr(conversion, 'convert_docx_to_pdf', lambda content, suffix: _pdf(5))
    assert len(conversion.pptx_to_slide_images(b'deck', 'deck.pptx')) == 2


def test_no_image_when_libreoffice_is_unavailable(monkeypatch):
    def unavailable(content, suffix):
        raise SofficeUnavailable('no engine')
    monkeypatch.setattr(conversion, 'convert_docx_to_pdf', unavailable)
    assert conversion.pptx_to_slide_images(b'deck', 'deck.pptx') == []


@pytest.mark.skipif(not shutil.which('soffice'), reason='LibreOffice is not installed')
def test_a_real_presentation_gives_one_image_per_slide():
    presentation = Presentation()
    for number in range(3):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f'Slide {number + 1}'
    buffer = io.BytesIO()
    presentation.save(buffer)
    images = conversion.pptx_to_slide_images(buffer.getvalue(), 'deck.pptx')
    assert len(images) == 3 and all(_is_png(image) for image in images)


def test_docx_body_keeps_the_text_and_an_empty_document_gets_a_placeholder():
    document = Document()
    document.add_paragraph('Exigence de contrôle')
    buffer = io.BytesIO()
    document.save(buffer)
    assert 'Exigence de contrôle' in conversion.docx_to_html_body(buffer.getvalue())
    empty = io.BytesIO()
    Document().save(empty)
    assert 'Document appears empty.' in conversion.docx_to_html_body(empty.getvalue())


def _antiword(monkeypatch, outputs):
    """Replace subprocess.run: ``outputs`` maps whether the UTF-8 mapping is requested to (return code, stdout)."""
    calls = []

    def run(command, **kwargs):
        calls.append(command)
        code, stdout = outputs['-m' in command]
        return SimpleNamespace(returncode=code, stdout=stdout, stderr=b'mapping file missing')

    monkeypatch.setattr(conversion.subprocess, 'run', run)
    monkeypatch.setattr(conversion, '_antiword_available', lambda: True)
    return calls


def test_antiword_is_asked_for_utf8_output(monkeypatch):
    calls = _antiword(monkeypatch, {True: (0, 'Vérification'.encode('utf-8'))})
    assert conversion.doc_to_text(b'doc') == 'Vérification'
    assert calls[0][1:3] == ['-m', 'UTF-8.txt'] and len(calls) == 1


def test_antiword_without_the_utf8_mapping_is_read_as_windows_1252(monkeypatch):
    calls = _antiword(monkeypatch, {True: (1, b''), False: (0, 'Vérification'.encode('cp1252'))})
    assert conversion.doc_to_text(b'doc') == 'Vérification'
    assert '-m' not in calls[1] and len(calls) == 2


def test_olefile_text_is_bounded_and_logged(monkeypatch, caplog):
    class FakeOle:
        def __init__(self, stream):
            pass

        def exists(self, name):
            return name == 'WordDocument'

        def openstream(self, name):
            runs = [f'Paragraph number {number:04d}'.encode('ascii') for number in range(10)]
            return io.BytesIO(b'\x00'.join(runs))

        def close(self):
            pass

    monkeypatch.setattr(conversion.olefile, 'OleFileIO', FakeOle)
    monkeypatch.setattr(conversion, '_OLEFILE_MAX_RUNS', 4)
    with caplog.at_level('WARNING'):
        text = conversion._olefile_extract(b'doc', 'a.doc')
    assert len(text.split('\n')) == 4 and 'keeping the first 4' in caplog.text


def test_olefile_gives_nothing_for_a_file_that_is_not_ole():
    assert conversion._olefile_extract(b'not an ole file', 'a.doc') is None
