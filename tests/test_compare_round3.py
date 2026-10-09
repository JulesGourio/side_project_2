"""Third pass of the 2026-10-04 Compare audit (docs/compare_audit_2026-10.md, part F).

Coverage:
  - /compare/impact and /compare/summarize through FastAPI (stream shape, cache)
  - chunked analysis: a failed part leaves a valid JSON array
  - "unrelated documents" warning
  - DOCX tables without a header row stay positional
  - relocation detection still works with the single shared index
"""

import asyncio
import hashlib
import io
import json
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document
from fastapi.testclient import TestClient

from server.services.processors._diff_engines import (
    extraction_warnings,
    looks_like_table_header,
    paragraph_semantic_diff,
)
from server.services.processors.docx import _extract_docx


# --- Routes ---

@pytest.fixture(scope='module')
def app():
    with (
        patch('server.app.init_lakebase', new_callable=AsyncMock),
        patch('server.app.shutdown_lakebase', new_callable=AsyncMock),
    ):
        from server.app import app as fastapi_app
        yield fastapi_app


@pytest.fixture
def client(app, monkeypatch):
    from server.services.user import require_compare
    for key, value in {
        'COMPARE_ENABLED': 'true', 'COMPARE_ANALYSIS_ENDPOINT': 'analysis', 'COMPARE_IMPACT_ENDPOINT': 'judge',
        'COMPARE_IMPACT_INDEX': 'cat.sch.chunks_index', 'COMPARE_SUMMARY_ENDPOINT': 'summary',
        'DATABRICKS_HOST': 'https://example.test', 'DATABRICKS_TOKEN': 'token',
    }.items():
        monkeypatch.setenv(key, value)
    app.dependency_overrides[require_compare] = lambda: None
    with (
        patch('server.routers.compare.get_user_identity', new=AsyncMock(return_value={'user_id': 'u', 'workspace_id': 'w'})),
        patch('server.routers.compare.store_impact_request', new=AsyncMock(return_value=7)),
        TestClient(app) as c,
    ):
        yield c
    app.dependency_overrides.clear()


CHANGES = json.dumps([{'section': '4.3 Torque', 'type': 'Value changed', 'criticality': 'High',
                       'before': 'couple de serrage 35 Nm', 'after': 'couple de serrage 40 Nm', 'rationale': ''}])
DOC = {'ref': 'PR-100', 'status': 'impacted', 'change_ids': ['C1'], 'max_score': 0.9, 'passages': []}


async def _fake_search(**kwargs):
    yield {'type': 'plan', 'changes': kwargs['extracted']['changes'], 'source': 'structured', 'queries_used': 1,
           'queries_failed': 0, 'chunks_returned': 3, 'candidates': 1, 'excluded_refs': [], 'not_judged': []}
    yield {'type': 'document', 'document': DOC}
    yield {'type': 'done', 'usage': {'input_tokens': 10, 'output_tokens': 5, 'total_tokens': 15, 'cost_eur': 0.0}}


def _events(res):
    return [json.loads(line) for line in res.text.splitlines() if line.strip()]


def test_impact_streams_plan_documents_done_and_caches_under_the_fingerprint(client):
    store = AsyncMock()
    with (
        patch('server.routers.compare.run_impact_search', new=_fake_search),
        patch('server.routers.compare.get_cached_impact_result', new=AsyncMock(return_value=None)),
        patch('server.routers.compare.store_impact_cache', new=store),
    ):
        res = client.post('/api/compare/impact', json={
            'changes_text': CHANGES, 'old_file_hash': 'a' * 64, 'new_file_hash': 'b' * 64,
        })
    events = _events(res)
    assert [e['type'] for e in events] == ['plan', 'document', 'done']
    assert events[0]['changes'][0]['id'] == 'C1'
    old_hash, new_hash, version, result = store.await_args.args
    assert (old_hash, new_hash) == ('a' * 64, 'b' * 64)
    assert ':' in version                      # APP_VERSION + fingerprint, not APP_VERSION alone
    assert result['documents'] == [DOC]


def test_impact_replays_a_cached_result(client):
    cached = {'changes': [{'id': 'C1'}], 'source': 'structured', 'documents': [DOC], 'usage': {}, 'duration_s': 1.0}
    with (
        patch('server.routers.compare.get_cached_impact_result', new=AsyncMock(return_value=cached)),
        patch('server.routers.compare.run_impact_search', side_effect=AssertionError('must not search')),
    ):
        res = client.post('/api/compare/impact', json={
            'changes_text': CHANGES, 'old_file_hash': 'a' * 64, 'new_file_hash': 'b' * 64,
        })
    events = _events(res)
    assert [e['type'] for e in events] == ['plan', 'document', 'done']
    assert events[0]['cached'] is True


def test_impact_reports_a_search_failure_as_an_error_event(client):
    async def _boom(**kwargs):
        raise RuntimeError('All 1 Vector Search queries failed')
        yield  # pragma: no cover — makes this an async generator

    with (
        patch('server.routers.compare.run_impact_search', new=_boom),
        patch('server.routers.compare.get_cached_impact_result', new=AsyncMock(return_value=None)),
        patch('server.routers.compare.store_error', new=AsyncMock()),
    ):
        res = client.post('/api/compare/impact', json={'changes_text': CHANGES})
    events = _events(res)
    assert events[-1]['type'] == 'error' and 'Vector Search' in events[-1]['error']


def test_impact_without_changes_does_not_search(client):
    with patch('server.routers.compare.run_impact_search', side_effect=AssertionError('must not search')):
        res = client.post('/api/compare/impact', json={'changes_text': '**Aucun changement significatif détecté.**'})
    events = _events(res)
    assert events[0]['no_changes'] is True and events[-1]['type'] == 'done'


def _docx_bytes(text):
    d = Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_summarize_caches_under_the_hash_of_the_received_bytes(client):
    data = _docx_bytes('The operator shall record the inspection results in the quality register.')
    store = AsyncMock()
    summary = AsyncMock(return_value={'summary': 'A short summary.', 'truncated': False, 'usage': {}})
    with (
        patch('server.routers.compare.get_cached_summary', new=AsyncMock(return_value=None)) as lookup,
        patch('server.routers.compare.store_summary_cache', new=store),
        patch('server.routers.compare.summarize_text', new=summary),
    ):
        res = client.post('/api/compare/summarize',
                          files={'file': ('proc.docx', data, 'application/octet-stream')},
                          data={'file_hash': 'stale-client-hash'})
    assert res.status_code == 200 and res.json()['summary'] == 'A short summary.'
    expected = hashlib.sha256(data).hexdigest()
    assert lookup.await_args.args[0] == expected
    assert store.await_args.args[0] == expected
    assert 'quality register' in summary.await_args.args[3]


def test_summarize_document_without_text_is_a_normal_answer(client):
    with patch('server.routers.compare.get_cached_summary', new=AsyncMock(return_value=None)):
        res = client.post('/api/compare/summarize',
                          files={'file': ('empty.docx', _docx_bytes(''), 'application/octet-stream')})
    assert res.status_code == 200 and res.json()['no_content'] is True


# --- Chunked analysis ---

def test_chunked_structured_stream_stays_valid_json_when_a_part_fails(monkeypatch):
    from server.services import chunked_analysis as ca

    async def fake_stream(host, token, ep, messages, *args, **kwargs):
        i = int(messages[-1]['content'][0]['text'])
        if i == 1:
            yield f'data: {json.dumps({"type": "error", "error": "endpoint timeout"})}\n\n'
        else:
            payload = json.dumps({'type': 'response.output_text.delta', 'delta': f'[{{"part": {i}}}]'})
            yield f'data: {payload}\n\n'
        yield 'data: [DONE]\n\n'

    monkeypatch.setattr(ca, 'stream_analysis', fake_stream)
    parts = [[{'role': 'user', 'content': [{'type': 'text', 'text': str(i)}]}] for i in range(3)]

    async def collect():
        text, errors = '', []
        async for chunk in ca.stream_analysis_chunked('h', 't', 'ep', parts, 100, 0, 0.0, structured=True):
            if not chunk.startswith('data: ') or '[DONE]' in chunk:
                continue
            event = json.loads(chunk[6:])
            if event.get('type') == 'response.output_text.delta':
                text += event['delta']
            elif event.get('type') == 'error':
                errors.append(event['error'])
        return text, errors

    text, errors = asyncio.run(collect())
    assert json.loads(text) == [{'part': 0}]
    assert errors == ['endpoint timeout']


# --- Engine ---

def _paragraphs(prefix, n):
    return '\n'.join(f'{prefix} paragraph {i} describes requirement number {i} of the {prefix} procedure in detail. '
                     f'[Page {1 + i // 10}]' for i in range(n))


def test_unrelated_documents_raise_a_warning():
    warnings = extraction_warnings(_paragraphs('welding', 40), _paragraphs('painting', 40))
    assert len(warnings) == 1 and 'almost no paragraph in common' in warnings[0]


def test_two_revisions_of_one_document_raise_no_warning():
    old = _paragraphs('welding', 40)
    new = old.replace('requirement number 7 ', 'requirement number 70 ')
    assert extraction_warnings(old, new) == []


def test_short_documents_never_raise_the_unrelated_warning():
    assert extraction_warnings(_paragraphs('welding', 5), _paragraphs('painting', 5)) == []


@pytest.mark.parametrize('cells, expected', [
    (['Fastener', 'Min torque', 'Max torque', 'Tool'], True),
    (['Activité', 'Phase 1', 'Responsable', 'Délai'], True),      # one numbered column is fine
    (['Reference', 'QR-2040'], False),                            # key/value form
    (['M6 bolt', '10 Nm', '12 Nm', 'Wrench A'], False),           # a data row
    (['Activity', '', 'Inspector', 'Manager'], False),            # empty cell
])
def test_looks_like_table_header(cells, expected):
    assert looks_like_table_header(cells) is expected


def _docx_table(rows):
    d = Document()
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return '\n'.join(_extract_docx(buf.getvalue())[0])


def test_docx_key_value_table_is_positional_and_one_change_stays_one_entry():
    form = [['Reference', 'QR-2040'], ['Revision', 'A'], ['Owner', 'Quality department']]
    old = _docx_table(form)
    assert 'Revision | A [Page 1, Para 2]' in old.splitlines()
    new = _docx_table([['Reference', 'QR-2041'], ['Revision', 'A'], ['Owner', 'Quality department']])
    diff, _ = paragraph_semantic_diff(old, new, page_label='Page')
    entries = [l for l in diff.splitlines() if l and not l.startswith('##')]
    assert entries and all('Reference' in e for e in entries)  # the other rows are untouched


def test_docx_table_without_header_keeps_empty_cells_in_place():
    rows = [['M8 bolt', '22 Nm', '', 'Wrench B'], ['M10 bolt', '', '45 Nm', 'Wrench C']]
    lines = _docx_table(rows).splitlines()
    assert lines[0].startswith('M8 bolt | 22 Nm |  | Wrench B')
    assert lines[1].startswith('M10 bolt |  | 45 Nm | Wrench C')


def test_paragraph_absorbed_into_a_longer_one_is_relocated_not_removed():
    moved = ('Personnel performing the inspection shall hold a valid qualification issued by the employer '
             'in accordance with the written practice and renewed every five years.')
    old = (f'1. SCOPE [Page 1]\nThis procedure applies to all parts. [Page 1]\n{moved} [Page 1]\n'
           '2. RECORDS [Page 2]\nRecords are kept for ten years in the archive room of the plant. [Page 2]')
    new = ('1. SCOPE [Page 1]\nThis procedure applies to all parts. [Page 1]\n'
           '2. RECORDS [Page 2]\nRecords are kept for ten years in the archive room of the plant. [Page 2]\n'
           f'3. QUALIFICATION [Page 3]\nThe following rule is unchanged from the previous issue: {moved} '
           + 'It is now part of this chapter together with the training requirements. ' * 6 + '[Page 3]')
    diff, _ = paragraph_semantic_diff(old, new, page_label='Page')
    assert 'RELOCATED' in diff
    assert not any(line.startswith('REMOVED') for line in diff.splitlines())
