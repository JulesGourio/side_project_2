"""Excel processor: workbook extensions, text extraction, keyed diff of tabular sheets."""

import io
import zipfile

import openpyxl

from server.services.compare.processors import excel
from server.services.compare.processors.factory import SUPPORTED_EXTENSIONS, file_type_of, get_processor


def _workbook(rows) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Data'
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _as_xlsm(xlsx: bytes) -> bytes:
    """The same workbook with the macro-enabled content type of an .xlsm file."""
    source = zipfile.ZipFile(io.BytesIO(xlsx))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',
                                    b'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
            target.writestr(item, data)
    return buffer.getvalue()


def test_extensions_are_read_in_any_case_and_macro_workbooks_are_supported():
    assert file_type_of('Plan.PDF') == 'pdf' and file_type_of('Tableau.XLSM') == 'excel'
    assert '.xlsm' in SUPPORTED_EXTENSIONS
    assert type(get_processor('book.xlsm')).__name__ == 'ExcelProcessor'


def test_a_macro_enabled_workbook_is_read_like_an_xlsx():
    lines, summary = excel._extract_xlsx(_as_xlsm(_workbook([['Ref', 'Value'], ['A', 1]])))
    assert lines[0] == '[Sheet: Data]' and 'A\t1' in lines and 'Data' in summary


HEADER = ['[Sheet: S]', 'Ref\tValue\tRank']


def _diff(old_rows, new_rows):
    return excel._build_diff_keyed(HEADER + old_rows, HEADER + new_rows, 'old.xlsx', 'new.xlsx')


def test_added_removed_and_modified_rows_are_reported_by_key():
    report = _diff(['A\t1\t1', 'B\t2\t2', 'C\t3\t3'], ['A\t1\t1', 'B\t9\t2', 'D\t4\t3'])
    assert 'ADDED rows (1)' in report and '+ D | Value: 4' in report
    assert 'REMOVED rows (1)' in report and '- C | Value: 3' in report
    assert "~ B:" in report and "Value: '2' → '9'" in report


def test_a_column_recalculated_in_most_of_many_rows_is_left_out():
    old = [f'K{i}\tx\t{i}' for i in range(25)]
    new = [f'K{i}\tx\t{i + 100}' for i in range(25)]
    report = _diff(old, new)
    assert 'Auto-excluded columns' in report and '"Rank"' in report
    assert 'No significant changes detected' in report


def test_a_column_changed_in_a_few_rows_is_reported_not_excluded():
    old = [f'K{i}\tx\t{i}' for i in range(6)]
    new = [f'K{i}\tx\t{i + 100}' for i in range(6)]
    report = _diff(old, new)
    assert 'Auto-excluded' not in report and "Rank: '3' → '103'" in report


def test_a_column_added_at_the_end_and_a_changed_header_are_reported():
    report = excel._build_diff_keyed(
        ['[Sheet: S]', 'Ref\tValue', 'A\t1'], ['[Sheet: S]', 'Ref\tValue\tUnit', 'A\t1\tmm'], 'old', 'new')
    assert 'Header changed' in report and "Unit: '' → 'mm'" in report


def test_a_sheet_present_only_in_the_new_workbook_is_reported():
    report = excel._build_diff_keyed(
        HEADER + ['A\t1\t1'], HEADER + ['A\t1\t1', '[Sheet: Extra]', 'Ref\tValue', 'B\t2'], 'old', 'new')
    assert '[Sheet: Extra]: sheet added (1 rows)' in report


def test_a_sheet_without_data_rows_gives_a_report_instead_of_failing():
    assert 'Key column: Ref' in _diff([], ['A\t1\t1'])


def test_the_diff_falls_back_when_the_first_column_is_no_key():
    assert _diff(['A\t1\t1', 'A\t2\t2'], ['A\t1\t1']) is None                   # repeated key
    assert _diff(['1\ta\tb', '2\tc\td'], ['1\ta\tb', '2\tc\td', '3\te\tf']) is None  # row numbers


def test_a_removed_sheet_is_reported():
    report = excel._build_diff_keyed(HEADER + ['A\t1\t1'], ['[Sheet: Other]', 'Ref\tValue', 'A\t1'], 'a', 'b')
    assert '[Sheet: S]: sheet removed entirely' in report


def test_a_value_that_looks_like_a_formula_is_exported_as_text():
    from server.services.compare.export_helpers import _build_excel_bytes
    rows = [{'section': 's', 'type': 'MODIFIED', 'criticality': 'high', 'before': '=HYPERLINK("http://x","y")',
             'after': 'ok', 'rationale': 'r'}]
    book = openpyxl.load_workbook(io.BytesIO(_build_excel_bytes(rows, '')))
    cells = [c for row in book.active.iter_rows() for c in row if isinstance(c.value, str) and 'HYPERLINK' in c.value]
    assert cells and all(c.data_type == 's' for c in cells)


def test_a_formula_changed_with_the_same_cached_value_is_read_with_its_formula():
    from server.services.compare.processors.excel import _extract_xlsx
    old_lines, _ = _extract_xlsx(_workbook([['Name', 'Total'], ['A', '=SUM(1,2)']]))
    new_lines, _ = _extract_xlsx(_workbook([['Name', 'Total'], ['A', '=SUM(1,3)']]))
    assert old_lines[2] == 'A\t {=SUM(1,2)}' and new_lines[2] == 'A\t {=SUM(1,3)}'
    assert old_lines != new_lines


def test_a_plain_value_is_read_without_braces():
    from server.services.compare.processors.excel import _extract_xlsx
    lines, _ = _extract_xlsx(_workbook([['Name', 'Total'], ['A', 5]]))
    assert lines == ['[Sheet: Data]', 'Name\tTotal', 'A\t5']


def _sheet(header, rows):
    return ['[Sheet: S]', '\t'.join(header)] + ['\t'.join(str(c) for c in r) for r in rows]


def _thirty_rows(columns=('Value', 'Rank', 'Note')):
    return [[f'K{i:02d}'] + [f'{c}{i}' for c in columns] for i in range(30)]


def test_a_column_inserted_in_the_middle_does_not_shift_the_other_columns():
    old = _sheet(['Ref', 'Value', 'Rank', 'Note'], _thirty_rows())
    new_rows = [[r[0], r[1], f'new{n}', r[2], r[3]] for n, r in enumerate(_thirty_rows())]
    new_rows[4][3] = 'Rank-changed'
    new = _sheet(['Ref', 'Value', 'Extra', 'Rank', 'Note'], new_rows)
    report = excel._build_diff_keyed(old, new, 'old.xlsx', 'new.xlsx')
    assert 'Columns added: "Extra"' in report
    assert 'Auto-excluded' not in report
    assert "~ K04:" in report and "Rank: 'Rank4' → 'Rank-changed'" in report
    assert report.count('  ~ ') == 1


def test_a_removed_column_is_reported_and_the_others_still_compared():
    old = _sheet(['Ref', 'Value', 'Rank', 'Note'], _thirty_rows())
    new_rows = [[r[0], r[2], r[3]] for r in _thirty_rows()]
    new_rows[2][2] = 'Note-changed'
    new = _sheet(['Ref', 'Rank', 'Note'], new_rows)
    report = excel._build_diff_keyed(old, new, 'old.xlsx', 'new.xlsx')
    assert 'Columns removed: "Value"' in report and 'Auto-excluded' not in report
    assert "Note: 'Note2' → 'Note-changed'" in report and report.count('  ~ ') == 1


def test_a_genuinely_recalculated_column_is_still_excluded_after_alignment():
    old = _sheet(['Ref', 'Value', 'Rank'], _thirty_rows(('Value', 'Rank')))
    new_rows = [[r[0], f'x{n}', r[1], r[2]] for n, r in enumerate(_thirty_rows(('Value', 'Rank')))]
    new_rows = [[r[0], r[1], r[2]] for r in new_rows]
    new = _sheet(['Ref', 'Extra', 'Rank'], [[r[0], 'e', f'changed{n}'] for n, r in enumerate(new_rows)])
    report = excel._build_diff_keyed(old, new, 'old.xlsx', 'new.xlsx')
    assert 'Auto-excluded columns' in report and '"Rank"' in report


def test_identical_headers_keep_the_positional_comparison():
    old = _sheet(['Ref', 'Value'], [['A', 1], ['B', 2]])
    new = _sheet(['Ref', 'Value'], [['A', 1], ['B', 3]])
    report = excel._build_diff_keyed(old, new, 'old.xlsx', 'new.xlsx')
    assert "Value: '2' → '3'" in report and 'Columns added' not in report
