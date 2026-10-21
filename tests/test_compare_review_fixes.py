"""Impact extraction limits, unverified judge quotes, chunked analysis merging, and processor extraction details."""

import asyncio
import io
import json

import pytest

from server.services.compare import chunked_analysis as ca
from server.services.compare.impact_queries import changes_to_queries, extract_changes
from server.services.infra.vector_search import _aggregate_docs, _judged_doc


# --- impact extraction ---

def test_long_markdown_section_is_split_into_changes_covering_every_bullet():
    bullets = [f'- **Item {i}** requirement number {i} was changed' for i in range(1, 41)]
    report = '## 4. Requirements\n' + '\n'.join(bullets) + '\n## 5. Other\n- one change\n'
    out = extract_changes(report, max_text_chars=400)
    section_changes = [c for c in out['changes'] if c['section'] == '4. Requirements']
    assert len(section_changes) > 1
    assert all(len(c['text']) <= 400 for c in out['changes'])
    joined = ' '.join(c['text'] for c in section_changes)
    assert 'Item 39 requirement number 39' in joined and 'Item 40 requirement number 40' in joined
    assert [c['id'] for c in out['changes']] == [f'C{i}' for i in range(1, len(out['changes']) + 1)]
    assert out['changes'][-1]['section'] == '5. Other'


def test_short_markdown_section_stays_one_change():
    out = extract_changes('## A\n- x changed\n- y added\n## B\n- z\n', max_text_chars=2000)
    assert [c['text'] for c in out['changes']] == ['A: x changed y added', 'B: z']


def test_configured_query_limit_reaches_the_extraction():
    report = '## S\n' + '\n'.join(f'- change number {i} ' + 'w' * 50 for i in range(30))
    res = changes_to_queries(report, max_queries=50, max_query_chars=300)
    assert all(len(q['text']) <= 300 for q in res['queries'])
    assert len(res['changes']) > 1


# --- unverified judge quote ---

def _cand():
    return _aggregate_docs([{
        'IDDOC': 'D1', 'REF': 'MI-0815', 'division': 'AS', 'url': 'u', 'score': 0.9, 'chunk_id': '1-000001',
        'chunk_text': '[Source: MI-0815 | Title: T | Division: AS | Category: X | Date de diffusion: 2020-01-01]\n\n'
                      'Serrer au couple de 12 N.m.',
        '_change_ids': ['C1'],
    }])[0]


def _verdict(quote):
    return {'impacted': True, 'confidence': 'high', 'reason': 'r',
            'passages': [{'passage': 1, 'changes': ['C1'], 'quote': quote}]}


def test_invented_quote_downgrades_the_document_to_check():
    doc = _judged_doc(_cand(), _verdict('Serrer au couple de 99 N.m.'), '')
    assert doc['status'] == 'check' and doc['impacted'] is True
    assert doc['passages'][0]['highlight'] is None


def test_verbatim_quote_keeps_the_document_impacted():
    assert _judged_doc(_cand(), _verdict('couple de 12 N.m.'), '')['status'] == 'impacted'


# --- chunked analysis ---

def test_split_body_never_glues_a_heading_to_the_previous_line():
    body = '\n## S1\n' + '\n'.join(f'ADDED: line {i} ' + 'x' * 30 for i in range(10)) + '\n## S2\nADDED: tail'
    parts = ca._split_body(body, 120)
    assert len(parts) > 2
    for part in parts:
        for line in part.split('\n'):
            assert '## ' not in line or line.startswith('## ')


def _messages(body, extra=''):
    text = f'intro\n--- TEXT CHANGES ---\n{body}\n\n--- VISUAL CHANGES ---{extra}'
    return [{'role': 'system', 'content': 's'}, {'role': 'user', 'content': [{'type': 'text', 'text': text}]}]


def test_continuation_part_carries_the_last_heading():
    body = '## 4. Torque\n' + '\n'.join(f'MODIFIED: value {i} ' + 'y' * 40 for i in range(12))
    parts = ca.split_messages_for_chunking(_messages(body), 100, 200)
    assert parts and len(parts) > 1
    texts = [p[1]['content'][0]['text'] for p in parts]
    assert 'continues the section' not in texts[0]
    assert all('continues the section "## 4. Torque"' in t for t in texts[1:])


def test_part_opening_with_its_own_heading_carries_nothing():
    body = '## A\n' + 'ADDED: a ' * 30 + '\n## B\n' + 'ADDED: b ' * 30
    parts = ca.split_messages_for_chunking(_messages(body), 100, 300)
    assert parts and len(parts) == 2
    assert 'continues the section' not in parts[1][1]['content'][0]['text']


def test_excel_diff_is_chunked_with_its_frame_kept():
    diff = '\n'.join(f'+row {i} ' + 'z' * 40 for i in range(40))
    text = ('Excel spreadsheet comparison\n  Document A: a.xlsx\n\nCell-level unified diff:\n\n'
            f'```diff\n{diff}\n```')
    msgs = [{'role': 'system', 'content': 's'}, {'role': 'user', 'content': [{'type': 'text', 'text': text}]}]
    parts = ca.split_messages_for_chunking(msgs, 500, 600)
    assert parts and len(parts) > 2
    rebuilt = []
    for p in parts:
        t = p[1]['content'][0]['text']
        assert t.startswith('Excel spreadsheet comparison') and t.endswith('\n```')
        assert t.count('```') == 2
        rebuilt += [ln for ln in t.split('```diff\n')[1].removesuffix('\n```').split('\n')]
    assert rebuilt == diff.split('\n')


def _run_chunked(monkeypatch, texts):
    async def fake_run_part(host, token, endpoint, msgs, *args, **kwargs):
        i = msgs[0]['i']
        return {'text': texts[i], 'usage': {}, 'error': None, 'passthrough': []}

    monkeypatch.setattr(ca, '_run_part', fake_run_part)

    async def collect():
        out = ''
        async for chunk in ca.stream_analysis_chunked('h', 't', 'e', [[{'i': i}] for i in range(len(texts))],
                                                      100, 0, 0.0, structured=False):
            if chunk.startswith('data: {') and 'output_text.delta' in chunk:
                out += json.loads(chunk[6:])['delta']
        return out

    return asyncio.run(collect())


def test_no_change_parts_are_dropped_between_real_changes(monkeypatch):
    out = _run_chunked(monkeypatch, ['- ADDED: a', 'No significant changes detected.', '- ADDED: c'])
    assert 'No significant changes' not in out
    assert out == '- ADDED: a\n\n- ADDED: c'


def test_no_change_statement_kept_when_no_part_has_changes(monkeypatch):
    out = _run_chunked(monkeypatch, ['No significant changes detected.', 'No significant changes detected.'])
    assert out == 'No significant changes detected.'


# --- processors ---

def test_dedup_repeated_headers_needs_three_pages():
    from server.services.compare.processors._diff_engines.collapse import _dedup_repeated_headers

    def item(idx, page):
        return {'idx': idx, 'clean': 'confidential', 'page': page}

    near_same_page = [item(1, 1), item(2, 1)]
    assert len(_dedup_repeated_headers(near_same_page)) == 2
    two_pages = [item(1, 1), item(2, 2)]
    assert len(_dedup_repeated_headers(two_pages)) == 2
    three_pages = [item(1, 1), item(2, 2), item(3, 3)]
    assert len(_dedup_repeated_headers(three_pages)) == 1
    no_pages = [item(1, None), item(2, None)]
    assert len(_dedup_repeated_headers(no_pages)) == 1


def test_image_cap_drops_labels_of_dropped_images_and_warns():
    from server.services.compare.processors._diff_engines.images import _MAX_IMAGE_BLOCKS, image_diff_blocks_dual
    from server.services.compare.processors.factory import diff_truncation_warnings

    new = {format(i, '064b'): {'b64': 'AAAA', 'pages': [i + 1]} for i in range(_MAX_IMAGE_BLOCKS + 5)}
    blocks = image_diff_blocks_dual({}, new)
    images = [b for b in blocks if b['type'] == 'image_url']
    labels = [b['text'] for b in blocks if b['type'] == 'text' and b['text'].startswith('ADDED image')]
    assert len(images) == len(labels) == _MAX_IMAGE_BLOCKS
    assert 'omitted' in blocks[-1]['text']
    warnings = diff_truncation_warnings([{'role': 'user', 'content': blocks}])
    assert len(warnings) == 1 and 'images' in warnings[0]
    assert diff_truncation_warnings([{'role': 'user', 'content': [{'type': 'text', 'text': 'x'}]}]) == []


def test_large_png_and_jpeg_are_downscaled_keeping_aspect_ratio():
    from PIL import Image
    from server.services.compare.processors.image import _normalise_image

    for ext, fmt, mime in (('.png', 'PNG', 'image/png'), ('.jpg', 'JPEG', 'image/jpeg')):
        buf = io.BytesIO()
        Image.new('RGB', (4000, 1000), 'white').save(buf, format=fmt)
        data, out_mime = _normalise_image(buf.getvalue(), ext)
        assert out_mime == mime
        assert Image.open(io.BytesIO(data)).size == (2000, 500)


def test_small_image_is_sent_unchanged():
    from PIL import Image
    from server.services.compare.processors.image import _normalise_image

    buf = io.BytesIO()
    Image.new('RGB', (100, 50), 'white').save(buf, format='PNG')
    assert _normalise_image(buf.getvalue(), '.png') == (buf.getvalue(), 'image/png')


def test_ppt_without_libreoffice_is_a_clear_error(monkeypatch):
    from server.services.compare import conversion
    from server.services.compare.processors.pptx import _extract_pptx
    from server.services.compare.soffice import SofficeUnavailable

    def none_found():
        raise SofficeUnavailable('absent')

    monkeypatch.setattr(conversion, 'find_soffice', none_found)
    with pytest.raises(ValueError, match=r'\.ppt format is not supported'):
        _extract_pptx(b'binary', 'old.ppt')


def test_ppt_is_converted_before_extraction(monkeypatch):
    from pptx import Presentation
    from server.services.compare.processors import pptx as pptx_mod

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = 'Hello'
    buf = io.BytesIO()
    prs.save(buf)
    calls = []
    monkeypatch.setattr(pptx_mod, 'ppt_to_pptx', lambda data, name: calls.append(name) or buf.getvalue())
    lines, _ = pptx_mod._extract_pptx(b'legacy', 'deck.PPT')
    assert calls == ['deck.PPT'] and lines == ['Hello [Slide 1]']


def test_pptx_speaker_notes_are_extracted_with_a_tag():
    from pptx import Presentation
    from server.services.compare.processors.pptx import _extract_pptx

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = 'Title'
    slide.notes_slide.notes_text_frame.text = 'Say the torque is 40'
    buf = io.BytesIO()
    prs.save(buf)
    lines, _ = _extract_pptx(buf.getvalue())
    assert '[Notes slide 1] Say the torque is 40 [Slide 1]' in '\n'.join(lines)


def test_docx_headers_and_footers_are_extracted_once():
    from docx import Document
    from server.services.compare.processors.docx.extract import _extract_docx

    doc = Document()
    doc.add_paragraph('Body text')
    doc.sections[0].header.paragraphs[0].text = 'ACME Confidential'
    doc.sections[0].footer.paragraphs[0].text = 'Rev B'
    doc.add_section()
    buf = io.BytesIO()
    doc.save(buf)
    lines, _ = _extract_docx(buf.getvalue())
    text = '\n'.join(lines)
    assert text.count('[Header] ACME Confidential') == 1
    assert text.count('[Footer] Rev B') == 1
