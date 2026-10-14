"""Golden messages of the PDF, DOCX, PPTX and XML processors on small synthetic file pairs.

tests/golden/processors.json holds the expected output; the LLM
messages, image pairs, warnings and metadata must stay byte-identical for every method.
"""

import hashlib
import io
import json
import random
from pathlib import Path

from server.services.compare.processors.docx import DocxProcessor
from server.services.compare.processors.pdf import PDFProcessor
from server.services.compare.processors.pptx import PptxProcessor
from server.services.compare.processors.xml import XMLProcessor

GOLDEN = Path(__file__).parent / 'golden' / 'processors.json'

_OLD = [
    '1. SCOPE',
    'The operator shall record every torque value measured on the primary fastener assembly before delivery.',
    'The supplier must deliver the part within ten working days after approval of the drawing by the customer.',
    '2. INSPECTION',
    'Inspection frequency is one part in ten for the first batch and one part in fifty afterwards.',
    'Records are kept for five years in the quality register of the plant concerned.',
]
_NEW = [
    '1. SCOPE',
    'The operator shall record every torque value measured on the primary fastener assembly before shipment.',
    'The supplier may deliver the part within ten working days after approval of the drawing by the customer.',
    '2. INSPECTION',
    'Inspection frequency is one part in ten for every batch.',
    'Records are kept for five years in the quality register of the plant concerned.',
    'A new paragraph states that the inspector signs each page of the traveller before release.',
]


def _png(seed: int, edit: bool = False) -> bytes:
    from PIL import Image, ImageDraw
    rng = random.Random(seed)
    img = Image.new('RGB', (120, 90), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    for _ in range(6):
        x0, y0 = rng.randrange(80), rng.randrange(60)
        draw.rectangle([x0, y0, x0 + rng.randrange(10, 40), y0 + rng.randrange(10, 30)], fill=(rng.randrange(256),) * 3)
    if edit:
        draw.line([0, 0, 119, 89], fill=(0, 0, 0), width=4)
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    return buf.getvalue()


def _pdf(lines, edit):
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for line in lines:
        page.insert_text((72, y), line[:95], fontsize=10)
        y += 18
    page.insert_image(pymupdf.Rect(72, 300, 192, 390), stream=_png(1, edit))
    return doc.tobytes()


def _docx(lines, edit):
    from docx import Document
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    table = doc.add_table(rows=2, cols=3)
    for c, text in enumerate(('Activity', 'Owner', 'Delay')):
        table.cell(0, c).text = text
    for c, text in enumerate(('Release', 'Quality', '5 days' if edit else '3 days')):
        table.cell(1, c).text = text
    doc.add_picture(io.BytesIO(_png(1, edit)))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx(lines, edit):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    for chunk in (lines[:3], lines[3:]):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = chunk[0]
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(3))
        box.text_frame.text = '\n'.join(chunk[1:])
    slide.shapes.add_picture(io.BytesIO(_png(1, edit)), Inches(1), Inches(5), Inches(2), Inches(1.5))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _xml(lines, edit):
    body = ''.join(f'<item id="{i}"><name>{line}</name></item>' for i, line in enumerate(lines))
    return f'<?xml version="1.0"?><root>{body}</root>'.encode()


def _scrub(messages):
    out = []
    for m in messages:
        content = m['content']
        if isinstance(content, list):
            content = [
                {'type': 'image', 'sha': hashlib.sha256(b['image_url']['url'].encode()).hexdigest()[:12]}
                if b.get('type') == 'image_url' else b for b in content
            ]
        out.append({'role': m['role'], 'content': content})
    return out


def _snapshot() -> dict:
    cases = {
        'pdf': (PDFProcessor, _pdf, 'a.pdf', 'b.pdf'),
        'docx': (DocxProcessor, _docx, 'a.docx', 'b.docx'),
        'pptx': (PptxProcessor, _pptx, 'a.pptx', 'b.pptx'),
        'xml': (XMLProcessor, _xml, 'a.xml', 'b.xml'),
    }
    out: dict = {}
    for name, (cls, make, old_name, new_name) in cases.items():
        old, new = make(_OLD, False), make(_NEW, True)
        for method in ('standard', 'structured', 'comparative'):
            for prompt in ('', 'custom system prompt'):
                r = cls(method).build_messages(old, old_name, new, new_name, prompt)
                out[f'{name}_{method}_{bool(prompt)}'] = {
                    'messages': _scrub(r.messages),
                    'metadata': [r.metadata.file_type, r.metadata.method, r.metadata.old_name, r.metadata.new_name],
                    'image_pairs': [
                        {k: (hashlib.sha256(v.encode()).hexdigest()[:12] if k.endswith('b64') and v else v)
                         for k, v in p.items()} for p in r.image_pairs
                    ],
                    'warnings': r.warnings,
                }
    return json.loads(json.dumps(out))


def test_processor_messages_match_golden():
    expected = json.loads(GOLDEN.read_text(encoding='utf-8'))
    actual = _snapshot()
    assert actual.keys() == expected.keys()
    mismatched = [k for k in expected if actual[k] != expected[k]]
    assert not mismatched, f'processor output changed for: {mismatched}'
