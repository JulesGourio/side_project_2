"""Preview endpoint — converts DOCX/PPTX/Excel to HTML for in-browser rendering."""

import asyncio
import base64
import io

import mammoth
import openpyxl
from fastapi import APIRouter, File, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, Response
from PIL import Image as PILImage
from pptx import Presentation

from ..services.conversion import doc_to_text, docx_to_html_body, pptx_to_slide_images
from ..services.soffice import SofficeUnavailable, convert_docx_to_pdf

router = APIRouter()

try:
    from pptx.enum.text import PP_ALIGN as _PP_ALIGN
    _ALIGN_MAP: dict = {_PP_ALIGN.CENTER: 'center', _PP_ALIGN.RIGHT: 'right', _PP_ALIGN.JUSTIFY: 'justify'}
except ImportError:
    _ALIGN_MAP = {}


@router.post('/preview', response_class=HTMLResponse)
async def preview_document(file: UploadFile = File(...)) -> HTMLResponse:
    content = await file.read()
    mime = file.content_type or ''
    name = file.filename or ''
    name_lower = name.lower()

    if 'wordprocessingml' in mime or name_lower.endswith('.docx'):
        return _preview_docx(content, name)

    if name_lower.endswith('.doc'):
        return _preview_doc(content, name)

    if 'presentationml' in mime or name_lower.endswith(('.pptx', '.ppt')):
        return _preview_pptx(content, name)

    if 'spreadsheet' in mime or 'excel' in mime or name_lower.endswith(('.xlsx', '.xls')):
        return _preview_excel(content, name)

    return HTMLResponse('<p>Preview not available for this file type.</p>', status_code=400)


@router.post('/preview/pdf')
async def preview_document_pdf(file: UploadFile = File(...)) -> Response:
    """Exact-layout preview: convert an office document (docx/doc/pptx/xlsx)
    to PDF through headless LibreOffice (server/services/soffice.py). Returns
    raw PDF bytes, or 501 when no engine is provisioned so the client can fall
    back to the HTML/docx-preview rendering."""
    content = await file.read()
    name = file.filename or 'document.docx'
    suffix = '.' + name.rsplit('.', 1)[-1].lower() if '.' in name else '.docx'
    try:
        pdf_bytes = await asyncio.to_thread(convert_docx_to_pdf, content, suffix=suffix)
    except SofficeUnavailable as e:
        return JSONResponse({'error': f'Exact PDF rendering unavailable: {e}'}, status_code=501)
    except Exception as e:
        return JSONResponse({'error': f'PDF conversion failed: {e}'}, status_code=502)
    return Response(content=pdf_bytes, media_type='application/pdf')


def _preview_docx(content: bytes, filename: str) -> HTMLResponse:
    return HTMLResponse(_wrap(docx_to_html_body(content), filename))


def _preview_doc(content: bytes, filename: str) -> HTMLResponse:
    # Some .doc files are actually valid OOXML — try mammoth first
    result = mammoth.convert_to_html(io.BytesIO(content))
    if result.value and len(result.value.strip()) > 20:
        return HTMLResponse(_wrap(result.value, filename))

    text = doc_to_text(content, filename)
    if text:
        paragraphs = '\n'.join(f'<p>{_esc(p)}</p>' for p in text.splitlines() if p.strip())
        return HTMLResponse(_wrap(paragraphs or '<p class="empty">Document appears empty.</p>', filename))

    return HTMLResponse(_wrap('<p class="empty">Preview could not be generated for this file.</p>', filename))


def _preview_pptx(content: bytes, filename: str) -> HTMLResponse:
    # Try LibreOffice headless PNG rendering first (pixel-perfect)
    slide_pngs = pptx_to_slide_images(content, filename)
    if slide_pngs:
        total = len(slide_pngs)
        slides_html: list[str] = []
        for idx, png_bytes in enumerate(slide_pngs, 1):
            b64 = base64.b64encode(png_bytes).decode('ascii')
            slides_html.append(
                f'<div class="slide-lo">'
                f'<span class="sl-num">{idx}<span class="sl-total"> / {total}</span></span>'
                f'<img src="data:image/png;base64,{b64}" alt="Slide {idx}">'
                f'</div>'
            )
        css_lo = (
            '.slide-lo{position:relative;margin-bottom:1.25rem;border-radius:10px;'
            'overflow:hidden;box-shadow:0 2px 10px rgba(0,0,0,.12);background:#fff;}'
            '.slide-lo img{width:100%;height:auto;display:block;}'
        )
        return HTMLResponse(_wrap('\n'.join(slides_html), filename, extra_css=css_lo))

    # Fallback: python-pptx positioned renderer
    prs = Presentation(io.BytesIO(content))
    total = len(prs.slides)
    slide_w = prs.slide_width or 9144000
    slide_h = prs.slide_height or 6858000
    aspect = round(slide_h / slide_w * 100, 3)

    slides_html = []
    for idx, slide in enumerate(prs.slides, 1):
        elements: list[tuple[int, str]] = []
        _pptx_collect(slide.shapes, slide_w, slide_h, elements)
        elements.sort(key=lambda x: x[0])

        inner = '\n'.join(h for _, h in elements) or '<span class="empty">No content</span>'
        slides_html.append(
            f'<div class="slide">'
            f'<span class="sl-num">{idx}<span class="sl-total"> / {total}</span></span>'
            f'<div class="slide-ratio" style="padding-top:{aspect}%">'
            f'<div class="slide-canvas">{inner}</div>'
            f'</div></div>'
        )

    css_pptx = (
        '.slide{position:relative;margin-bottom:1.5rem;border-radius:10px;overflow:hidden;'
        'box-shadow:0 2px 10px rgba(0,0,0,.1);background:#fff;}'
        '.slide-ratio{position:relative;width:100%;}'
        '.slide-canvas{position:absolute;inset:0;overflow:hidden;}'
        '.sl-pos{position:absolute;box-sizing:border-box;overflow:hidden;}'
        '.sl-pos p{margin:0 0 .1em;line-height:1.25;white-space:pre-wrap;word-break:break-word;}'
        '.sl-pos table{border-collapse:collapse;width:100%;font-size:.75em;}'
        '.sl-pos td,.sl-pos th{border:1px solid rgba(0,0,0,.18);padding:.15em .3em;}'
        '.sl-pos th{background:rgba(0,0,0,.07);font-weight:600;}'
        '.sl-pos img{width:100%;height:100%;object-fit:contain;display:block;}'
    )
    return HTMLResponse(_wrap('\n'.join(slides_html), filename, extra_css=css_pptx))


def _pptx_collect(shapes, slide_w: int, slide_h: int, out: list) -> None:
    """Recursively collect positioned HTML elements from a slide's shapes."""
    for shape in shapes:
        try:
            left  = (getattr(shape, 'left',   0) or 0)
            top   = (getattr(shape, 'top',    0) or 0)
            width = (getattr(shape, 'width',  0) or 0)
            height= (getattr(shape, 'height', 0) or 0)
            lp = left   / slide_w * 100
            tp = top    / slide_h * 100
            wp = width  / slide_w * 100
            hp = height / slide_h * 100
            pos = f'left:{lp:.3f}%;top:{tp:.3f}%;width:{wp:.3f}%;height:{hp:.3f}%;'
        except Exception:
            pos = 'left:0;top:0;width:100%;height:auto;'
            top = 0

        try:
            if shape.shape_type == 6:  # GROUP
                _pptx_collect(shape.shapes, slide_w, slide_h, out)
                continue
        except Exception:
            pass

        try:
            if shape.has_table:
                rows = []
                for r_idx, row in enumerate(shape.table.rows):
                    cells = []
                    for cell in row.cells:
                        tag = 'th' if r_idx == 0 else 'td'
                        try:
                            txt = cell.text_frame.text
                        except Exception:
                            txt = ''
                        cells.append(f'<{tag}>{_esc(txt)}</{tag}>')
                    rows.append(f'<tr>{"".join(cells)}</tr>')
                html = f'<div class="sl-pos" style="{pos}"><table>{"".join(rows)}</table></div>'
                out.append((top, html))
                continue
        except Exception:
            pass

        try:
            if shape.shape_type in (13, 14):
                img_b = _normalise_image(shape.image.blob, max_px=1200)
                b64 = base64.b64encode(img_b).decode('ascii')
                html = f'<div class="sl-pos" style="{pos}"><img src="data:image/jpeg;base64,{b64}" alt=""></div>'
                out.append((top, html))
        except Exception:
            pass

        try:
            if not (hasattr(shape, 'has_text_frame') and shape.has_text_frame):
                continue
            paras = []
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                pstyle = ''
                try:
                    sz = para.runs[0].font.size
                    if sz:
                        pt = max(6, min(96, sz / 12700))
                        pstyle += f'font-size:{pt:.0f}pt;'
                except Exception:
                    pass
                try:
                    if para.runs[0].font.bold:
                        pstyle += 'font-weight:700;'
                except Exception:
                    pass
                try:
                    align = para.alignment
                    if align:
                        a = _ALIGN_MAP.get(align)
                        if a:
                            pstyle += f'text-align:{a};'
                except Exception:
                    pass
                paras.append(f'<p style="{pstyle}">{_esc(text)}</p>')
            if paras:
                html = f'<div class="sl-pos" style="{pos}">{"".join(paras)}</div>'
                out.append((top, html))
        except Exception:
            pass


def _preview_excel(content: bytes, filename: str) -> HTMLResponse:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        return HTMLResponse(_wrap(f'<p class="warning">Could not open Excel file: {_esc(str(e))}</p>', filename))

    sheets_html: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_html: list[str] = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx >= 500:
                rows_html.append('<tr><td colspan="999" class="more-rows">… (preview limited to 500 rows)</td></tr>')
                break
            cells = ''.join(
                f'<{"th" if row_idx == 0 else "td"}>{_esc(str(c or ""))}</{"th" if row_idx == 0 else "td"}>'
                for c in row
            )
            rows_html.append(f'<tr>{cells}</tr>')
        sheets_html.append(
            f'<section class="sheet"><h2>{_esc(sheet_name)}</h2>'
            f'<div class="table-wrap"><table>{"".join(rows_html)}</table></div></section>'
        )

    wb.close()
    return HTMLResponse(_wrap(''.join(sheets_html), filename))


def _normalise_image(img_bytes: bytes, max_px: int = 960) -> bytes:
    img = PILImage.open(io.BytesIO(img_bytes)).convert('RGB')
    w, h = img.size
    scale = min(1.0, max_px / max(w, h, 1))
    if scale < 1.0:
        img = img.resize((max(1, int(w * scale)), max(1, int(h * scale))), PILImage.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format='JPEG', quality=85)
    return buf.getvalue()


def _esc(text: str) -> str:
    return (
        text.replace('&', '&amp;')
            .replace('<', '&lt;')
            .replace('>', '&gt;')
    )


def _wrap(body: str, filename: str, extra_css: str = '') -> str:
    return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>
  *, *::before, *::after {{ box-sizing: border-box; }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
    margin: 0; padding: 1.25rem 1.75rem 2rem;
    color: #1f2937; background: #f8fafc; font-size: 14px; line-height: 1.6;
  }}
  h1, h2, h3, h4 {{ color: #111827; margin: 1.1em 0 .35em; font-weight: 600; }}
  p {{ margin: .4em 0; }}
  ul, ol {{ padding-left: 1.4em; margin: .4em 0; }}
  img {{ max-width: 100%; height: auto; border-radius: 4px; }}
  .table-wrap {{ overflow: auto; max-height: 560px; border: 1px solid #e2e8f0; border-radius: 8px; }}
  .table-wrap table {{ border-collapse: collapse; font-size: 13px; min-width: 100%; }}
  .table-wrap td, .table-wrap th {{ border: 1px solid #e2e8f0; padding: .3em .55em; white-space: nowrap; }}
  .table-wrap th {{ background: #f1f5f9; font-weight: 600; position: sticky; top: 0; z-index: 1; }}
  .table-wrap tr:hover td {{ background: #f8fafc; }}
  .sheet {{ margin-bottom: 2.5rem; }}
  .sheet h2 {{ font-size: .85rem; text-transform: uppercase; letter-spacing: .06em;
               color: #64748b; border-bottom: 1px solid #e2e8f0; padding-bottom: .4em; }}
  .sl-num {{
    position: absolute; top: 10px; right: 12px; z-index: 10;
    background: rgba(0,0,0,.50); color: #fff;
    font-size: .72rem; font-weight: 700; padding: 2px 8px; border-radius: 99px;
    letter-spacing: .04em; line-height: 1.6; pointer-events: none;
  }}
  .sl-total {{ font-weight: 400; opacity: .8; }}
  .warning {{
    background: #fef9c3; border: 1px solid #fde047; border-radius: 8px;
    padding: .8rem 1rem; color: #713f12; margin-bottom: 1rem; font-size: .88rem;
  }}
  .empty {{ color: #94a3b8; font-style: italic; font-size: .85rem; }}
  .more-rows {{ color: #94a3b8; text-align: center; padding: .5em; font-style: italic; }}
  {extra_css}
</style>
</head>
<body>
{body}
</body>
</html>"""
