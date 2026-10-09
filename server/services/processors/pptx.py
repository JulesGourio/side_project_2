"""PPTX / PPT processor — three processing methods.

standard:    paragraph semantic diff on slide text + focused Markdown bullet-list output.
structured:  paragraph semantic diff on slide text + strict JSON array output (for Excel export).
comparative: section canonical diff on slide text + section-grouped output.

Embedded images go through the same perceptual hashing pipeline as PDF/DOCX
(_compute_image_diff_pairs in _diff_engines), so PPTX now gets the same
MODIFIED / REMOVED / ADDED / ORIENTATION categorisation instead of a coarse
"set of MD5 hashes per slide" comparison.
"""

import base64
import io
import logging
from typing import Any, Dict, List, Tuple

from .base import BaseProcessor, ProcessMetadata, ProcessResult
from ._diff_engines import (
    SYSTEM_PROMPT_STANDARD,
    SYSTEM_PROMPT_STRUCTURED,
    dhash,
    extraction_warnings,
    hash16,
    image_diff_blocks_dual,
    image_diff_pairs,
    images_are_similar,
    paragraph_semantic_diff,
    section_canonical_diff,
    truncate_diff,
)

logger = logging.getLogger(__name__)

_IMG_MAX_PX = 1024
_IMG_JPEG_QUALITY = 80


def _shapes_text(shapes) -> List[str]:
    """Recursively extract text from shapes, including tables and groups."""
    texts: List[str] = []
    for shape in shapes:
        try:
            if shape.shape_type == 6:  # group
                texts.extend(_shapes_text(shape.shapes))
                continue
        except Exception:
            pass
        try:
            # Charts carry technical content (title, series, categories) that
            # would otherwise be invisible to the text diff.
            if getattr(shape, 'has_chart', False):
                chart = shape.chart
                parts: List[str] = []
                try:
                    if chart.has_title:
                        title = chart.chart_title.text_frame.text.strip()
                        if title:
                            parts.append(title)
                except Exception:
                    pass
                try:
                    names = [str(s.name) for s in chart.series if getattr(s, 'name', None)]
                    if names:
                        parts.append('series: ' + ', '.join(names))
                except Exception:
                    pass
                try:
                    cats = [str(c) for c in chart.plots[0].categories if c is not None]
                    if cats:
                        parts.append('categories: ' + ', '.join(cats[:20]))
                except Exception:
                    pass
                if parts:
                    texts.append('[Chart] ' + ' | '.join(parts))
                continue
        except Exception:
            pass
        try:
            if shape.has_table:
                for row in shape.table.rows:
                    row_cells = []
                    for cell in row.cells:
                        try:
                            t = cell.text_frame.text.strip()
                            if t:
                                row_cells.append(t)
                        except Exception:
                            pass
                    if row_cells:
                        texts.append(' | '.join(row_cells))
                continue
        except Exception:
            pass
        try:
            if hasattr(shape, 'has_text_frame') and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        texts.append(text)
        except Exception:
            pass
    return texts


def _collect_shape_images(shapes, slide_num: int, images: Dict[str, Dict[str, Any]]) -> None:
    """Recursively extract images from shapes, dedupe by perceptual dhash, record slides.

    Builds the same dict shape used by DOCX/PDF processors:
        { dhash_str: { 'b64': str, 'pages': [slide_nums] } }
    so the shared pairing engine in _diff_engines can be applied consistently.
    """
    from PIL import Image as PILImage
    for shape in shapes:
        try:
            if shape.shape_type == 6:  # group
                _collect_shape_images(shape.shapes, slide_num, images)
                continue
        except Exception:
            pass
        try:
            if shape.shape_type not in (13, 14):  # picture, placeholder picture
                continue
            img_bytes = shape.image.blob
            # Normalise to JPEG with consistent size and quality so the hash
            # remains stable across small re-saves of the same source image.
            img = PILImage.open(io.BytesIO(img_bytes)).convert('RGB')
            w, h = img.size
            scale = min(1.0, _IMG_MAX_PX / max(w, h, 1))
            if scale < 1.0:
                img = img.resize(
                    (max(1, int(w * scale)), max(1, int(h * scale))),
                    PILImage.Resampling.LANCZOS,
                )
            buf = io.BytesIO()
            img.save(buf, format='JPEG', quality=_IMG_JPEG_QUALITY)
            jpeg_bytes = buf.getvalue()

            h_key = dhash(jpeg_bytes)
            # Match against existing dhashes within a small Hamming budget so
            # near-duplicates (one slide reusing another's picture with tiny
            # re-compression) deduplicate to a single entry with multiple slides.
            matched = next((k for k in images if images_are_similar(k, h_key)), None)
            if matched:
                if slide_num not in images[matched]['pages']:
                    images[matched]['pages'].append(slide_num)
            else:
                b64 = base64.b64encode(jpeg_bytes).decode('ascii')
                images[h_key] = {'b64': b64, 'pages': [slide_num], 'hash16': hash16(jpeg_bytes)}
        except Exception as e:
            logger.debug('PPTX image extraction skipped (slide %d): %s', slide_num, e)


def _extract_pptx(pptx_bytes: bytes) -> Tuple[List[str], Dict[str, Dict[str, Any]]]:
    """Return (slide_text_lines, images_dict).

    Each text line is tagged with a trailing [Slide N] tag so the diff engine
    can track which slide a change originates from. Images are returned in the
    same shape used by DOCX/PDF (see _collect_shape_images).
    """
    from pptx import Presentation
    prs = Presentation(io.BytesIO(pptx_bytes))
    slide_texts: List[str] = []
    images: Dict[str, Dict[str, Any]] = {}
    for slide_num, slide in enumerate(prs.slides, start=1):
        lines = [f'{t} [Slide {slide_num}]' for t in _shapes_text(slide.shapes)]
        slide_texts.append('\n'.join(lines) if lines else f'[Slide {slide_num}]')
        _collect_shape_images(slide.shapes, slide_num, images)
    return slide_texts, images


# --- Processor ---

class PptxProcessor(BaseProcessor):

    def __init__(self, method: str = 'standard') -> None:
        self.method = method

    def build_messages(
        self,
        old_bytes: bytes,
        old_name: str,
        new_bytes: bytes,
        new_name: str,
        system_prompt: str = '',
    ) -> ProcessResult:
        if self.method == 'structured':
            return self._paragraph_diff(
                old_bytes, old_name, new_bytes, new_name,
                SYSTEM_PROMPT_STRUCTURED, 'structured',
            )
        if self.method == 'comparative':
            return self._comparative_diff(old_bytes, old_name, new_bytes, new_name, system_prompt)
        return self._paragraph_diff(
            old_bytes, old_name, new_bytes, new_name,
            SYSTEM_PROMPT_STANDARD, 'standard',
        )

    def _paragraph_diff(
        self,
        old_bytes: bytes,
        old_name: str,
        new_bytes: bytes,
        new_name: str,
        system_prompt: str,
        method_name: str,
    ) -> ProcessResult:
        logger.info('PptxProcessor [%s]: paragraph semantic diff — %s → %s',
                    method_name, old_name, new_name)
        old_slides, old_imgs = _extract_pptx(old_bytes)
        new_slides, new_imgs = _extract_pptx(new_bytes)

        old_text = '\n'.join(old_slides)
        new_text = '\n'.join(new_slides)
        warnings = extraction_warnings(old_text, new_text)
        diff_text, filtered = paragraph_semantic_diff(old_text, new_text, page_label='Slide')
        diff_text = truncate_diff(diff_text)

        # Slides are ordered, so pass show_position=True to use the page-aware
        # pairing path (same as PDF). Slide numbers replace page numbers.
        img_blocks = image_diff_blocks_dual(old_imgs, new_imgs, show_position=True)
        img_pairs  = image_diff_pairs(old_imgs, new_imgs, show_position=True)

        intro = (
            f'Global paragraph alignment across slides: {filtered} trivial lines filtered.\n'
            'Similar text blocks shown as MODIFIED with ~~removed~~ and **added** words inline.\n'
            'Each block tagged with its source slide.\n\n'
            f'--- TEXT CHANGES ---\n{diff_text}\n\n--- VISUAL CHANGES ---'
        )
        if not diff_text.strip():
            intro = 'No significant text changes detected.\n\n--- VISUAL CHANGES ---'

        content_blocks: List[Dict[str, Any]] = [{'type': 'text', 'text': intro}] + img_blocks
        messages = [
            {'role': 'system', 'content': system_prompt},
            {'role': 'user', 'content': content_blocks},
        ]
        return ProcessResult(
            messages=messages,
            metadata=ProcessMetadata(
                file_type='pptx', method=method_name,
                old_name=old_name, new_name=new_name,
            ),
            image_pairs=img_pairs,
            warnings=warnings,
        )

    def _comparative_diff(
        self,
        old_bytes: bytes,
        old_name: str,
        new_bytes: bytes,
        new_name: str,
        system_prompt: str,
    ) -> ProcessResult:
        logger.info('PptxProcessor [comparative]: section canonical diff — %s → %s',
                    old_name, new_name)
        old_slides, old_imgs = _extract_pptx(old_bytes)
        new_slides, new_imgs = _extract_pptx(new_bytes)

        old_text = '\n'.join(old_slides)
        new_text = '\n'.join(new_slides)
        diff_text, filtered = section_canonical_diff(old_text, new_text)
        diff_text = truncate_diff(diff_text)

        img_blocks = image_diff_blocks_dual(old_imgs, new_imgs, show_position=True)
        img_pairs  = image_diff_pairs(old_imgs, new_imgs, show_position=True)

        intro = (
            f'Section canonical diff across slides: {filtered} trivial pairs filtered.\n\n'
            f'--- CHANGES BY SECTION ---\n{diff_text}\n\n--- VISUAL CHANGES ---'
        )
        if not diff_text.strip():
            intro = 'No significant text changes detected.\n\n--- VISUAL CHANGES ---'

        content_blocks: List[Dict[str, Any]] = [{'type': 'text', 'text': intro}] + img_blocks
        messages: List[Dict[str, Any]] = []
        if system_prompt:
            messages.append({'role': 'system', 'content': system_prompt})
        messages.append({'role': 'user', 'content': content_blocks})
        return ProcessResult(
            messages=messages,
            metadata=ProcessMetadata(
                file_type='pptx', method='comparative',
                old_name=old_name, new_name=new_name,
            ),
            image_pairs=img_pairs,
        )
