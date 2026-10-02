"""Pure export helper functions — no FastAPI / Databricks / asyncpg dependencies.

Extracted here so the notebook (debug_prompt.ipynb) can import them directly
without triggering the server-side import chain in routers/exports.py.
"""

import base64
import io
import json
import logging
import re
from typing import Any

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# JSON parsing
# ---------------------------------------------------------------------------

def _parse_json_response(text: str):
    """Robust JSON array parser — three fallback strategies."""
    stripped = text.strip()
    try:
        data = json.loads(stripped)
        if isinstance(data, list):
            return data
    except json.JSONDecodeError:
        pass

    m = re.search(r'`{3}(?:json)?\s*([\s\S]*?)`{3}', stripped)
    if m:
        try:
            data = json.loads(m.group(1).strip())
            if isinstance(data, list):
                return data
        except json.JSONDecodeError:
            pass

    m = re.search(r'(\[[\s\S]*\])', stripped)
    if m:
        try:
            data = json.loads(m.group(1))
            if isinstance(data, list):
                return data
        except json.JSONDecodeError:
            pass
    return None


# ---------------------------------------------------------------------------
# Image helpers
# ---------------------------------------------------------------------------

def _make_composite_image(old_b64: str | None, new_b64: str | None, width: int = 200, height: int = 80):
    """Return a BytesIO of a side-by-side OLD|NEW JPEG thumbnail, or None on failure."""
    try:
        from PIL import Image as PILImage, ImageDraw
        composite = PILImage.new('RGB', (width, height), color=(255, 255, 255))
        half = width // 2
        if old_b64:
            old_img = PILImage.open(io.BytesIO(base64.b64decode(old_b64))).convert('RGB')
            old_img.thumbnail((half - 2, height - 2), PILImage.Resampling.LANCZOS)
            composite.paste(old_img, (1, (height - old_img.height) // 2))
        if new_b64:
            new_img = PILImage.open(io.BytesIO(base64.b64decode(new_b64))).convert('RGB')
            new_img.thumbnail((half - 2, height - 2), PILImage.Resampling.LANCZOS)
            composite.paste(new_img, (half + 1, (height - new_img.height) // 2))
        draw = ImageDraw.Draw(composite)
        draw.line([(half, 0), (half, height)], fill=(200, 200, 200), width=1)
        buf = io.BytesIO()
        composite.save(buf, format='JPEG', quality=80)
        buf.seek(0)
        return buf
    except Exception as e:
        logger.debug('Composite image failed: %s', e)
        return None


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def _build_excel_bytes(rows: list, ft: str, image_pairs: list | None = None) -> bytes:
    """Generate a colour-coded Excel workbook from a list of diff items. ft = file type."""
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ft = (ft or '').lower().strip()

    if ft in ('pdf', 'docx'):
        _page_key = 'page'
        _KEYS = ['section', 'page', 'type', 'criticality', 'before', 'after', 'rationale']
        _HDRS = ['Section', 'Page', 'Type', 'Criticality', 'Before', 'After', 'Rationale']
        _COL_WIDTHS = [22, 8, 24, 16, 42, 42, 50]
        _page_col_idx = 2
    elif ft == 'pptx':
        _page_key = 'page'
        _KEYS = ['section', 'page', 'type', 'criticality', 'before', 'after', 'rationale']
        _HDRS = ['Section', 'Slide', 'Type', 'Criticality', 'Before', 'After', 'Rationale']
        _COL_WIDTHS = [22, 8, 24, 16, 42, 42, 50]
        _page_col_idx = 2
    elif ft == 'xml':
        _page_key = 'page'
        _KEYS = ['section', 'page', 'type', 'criticality', 'before', 'after', 'rationale']
        _HDRS = ['Section', 'No.', 'Type', 'Criticality', 'Before', 'After', 'Rationale']
        _COL_WIDTHS = [22, 6, 24, 16, 42, 42, 50]
        _page_col_idx = None
    else:
        _page_key = None
        _KEYS = ['section', 'type', 'criticality', 'before', 'after', 'rationale']
        _HDRS = ['Section', 'Type', 'Criticality', 'Before', 'After', 'Rationale']
        _COL_WIDTHS = [22, 24, 16, 42, 42, 50]
        _page_col_idx = None

    has_images = bool(image_pairs)
    if has_images:
        _HDRS = list(_HDRS) + ['Image']
        _COL_WIDTHS = list(_COL_WIDTHS) + [30]

    _CRIT_COLORS = {'high': 'F8CECC', 'medium': 'FFF2CC', 'low': 'D5E8D4'}
    _TYPE_COLORS = {
        'value changed':        'FFF2CC',
        'reference updated':    'DAE8FC',
        'requirement modified': 'D5E8D4',
        'requirement added':    'C6EFCE',
        'requirement removed':  'F8CECC',
        'procedure changed':    'E1D5E7',
        'scope changed':        'FFE6CC',
        'modal change':         'E8D5F5',
    }
    _IMAGE_KEYWORDS = {
        'image', 'figure', 'visual', 'photo', 'diagram', 'illustration',
        'screenshot', 'picture', 'schematic', 'drawing', 'graphic', 'chart',
        'plot',
    }

    def _is_visual_row(item: dict) -> bool:
        t = str(item.get('type', '') or '').lower()
        return any(kw in t for kw in _IMAGE_KEYWORDS)

    def _row_page_int(item: dict):
        """First page number in the cell — used for image-pair matching (old page)."""
        raw = str(item.get('page', '') or '') if _page_key else ''
        m = re.search(r'\d+', raw)
        return int(m.group()) if m else None

    def _row_page_new_int(item: dict):
        """Last page number in the cell — the NEW-document page for "old→new" entries.
        The report describes the new document, so rows are ordered by where the
        reader finds them there."""
        raw = str(item.get('page', '') or '') if _page_key else ''
        nums = re.findall(r'\d+', raw)
        return int(nums[-1]) if nums else None

    def _section_tuple(item: dict):
        """Parse a dotted section number ("6.3.1" -> (6, 3, 1)) for stable
        secondary ordering. Returns None for named/unnumbered sections
        (TABLE OF CONTENTS, APPENDIX A, ...) so they sort after numbered ones
        on the same page."""
        sec = str(item.get('section', '') or '')
        m = re.match(r'\s*(\d+(?:\.\d+)+)', sec)
        return tuple(int(x) for x in m.group(1).split('.')) if m else None

    _BIG = 10 ** 9

    def _page_sort_key(item):
        if not isinstance(item, dict) or not _page_key:
            return (1, _BIG, 1, ())
        p = _row_page_new_int(item)
        page_part = (0, p) if p is not None else (1, _BIG)
        sec = _section_tuple(item)
        sec_part = (0, sec) if sec is not None else (1, ())
        return (*page_part, *sec_part)

    img_pairs_list: list = list(image_pairs or [])
    if has_images and img_pairs_list:
        visual_rows = [it for it in rows if isinstance(it, dict) and _is_visual_row(it)]
        used_pair_idx: set = set()

        for item in visual_rows:
            page = _row_page_int(item)
            if page is None:
                continue
            for i, pair in enumerate(img_pairs_list):
                if i in used_pair_idx:
                    continue
                if pair.get('new_page') == page or pair.get('old_page') == page:
                    item['_img_pair'] = pair
                    used_pair_idx.add(i)
                    break

        remaining_pair_idx = [i for i in range(len(img_pairs_list)) if i not in used_pair_idx]
        remaining_rows = [it for it in visual_rows if '_img_pair' not in it]
        for item, pi in zip(remaining_rows, remaining_pair_idx):
            item['_img_pair'] = img_pairs_list[pi]

    if _page_key:
        rows = sorted(rows, key=_page_sort_key)

    img_col_num = len(_HDRS)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Changes'
    thin = Side(style='thin', color='CCCCCC')
    brd = Border(left=thin, right=thin, top=thin, bottom=thin)

    for ci, h in enumerate(_HDRS, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(bold=True, color='FFFFFF', size=11)
        c.fill = PatternFill('solid', fgColor='2F5496')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = brd
        if ft == 'docx' and _page_col_idx and ci == _page_col_idx:
            c.comment = Comment(
                'Page numbers are approximate (estimated from page-break markers in the DOCX file). '
                'The actual change may appear on an adjacent page.',
                'Qualibot',
            )
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'
    for ci in range(1, len(_HDRS) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = _COL_WIDTHS[ci - 1] if ci <= len(_COL_WIDTHS) else 30

    _crit_idx = _KEYS.index('criticality')
    _type_idx = _KEYS.index('type')
    _page_key_idx = _KEYS.index('page') if _page_key and 'page' in _KEYS else None

    for ri, item in enumerate(rows, 2):
        if not isinstance(item, dict):
            continue
        row = [str(item.get(k, '')) for k in _KEYS]
        crit_val = row[_crit_idx].lower()
        type_val = row[_type_idx].lower()
        fgcol = _CRIT_COLORS.get(crit_val) or _TYPE_COLORS.get(type_val, 'FFFFFF')
        fill = PatternFill('solid', fgColor=fgcol)
        for ci, val in enumerate(row, 1):
            cell_value: Any = val
            if _page_key_idx is not None and ci == _page_key_idx + 1:
                if '→' in val or '->' in val:
                    cell_value = val
                else:
                    m = re.search(r'\d+', val)
                    if m:
                        cell_value = int(m.group())
            c = ws.cell(row=ri, column=ci, value=cell_value)
            c.alignment = Alignment(
                horizontal='center' if (_page_key_idx is not None and ci == _page_key_idx + 1) else 'left',
                vertical='top',
                wrap_text=True,
            )
            c.fill = fill
            c.border = brd
        ws.row_dimensions[ri].height = 35

        if has_images:
            img_pair = item.get('_img_pair') if _is_visual_row(item) else None
            img_cell = ws.cell(row=ri, column=img_col_num)
            img_cell.fill = fill
            img_cell.border = brd

            if img_pair:
                composite = _make_composite_image(img_pair.get('old_b64'), img_pair.get('new_b64'))
                if composite:
                    try:
                        xl_img = XLImage(composite)
                        xl_img.width = 220
                        xl_img.height = 75
                        col_letter = get_column_letter(img_col_num)
                        ws.add_image(xl_img, f'{col_letter}{ri}')
                        ws.row_dimensions[ri].height = 60
                    except Exception as exc:
                        logger.warning('Excel image insert failed row %d: %s', ri, exc)
                        img_cell.value = f'[{img_pair.get("status", "?")}]'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF helpers
# ---------------------------------------------------------------------------

# fpdf2 core fonts are latin-1 only — map Unicode chars to safe latin-1 equivalents.
_UNICODE_TO_LATIN1 = str.maketrans({
    '—': '-', '–': '-', '’': "'", '‘': "'",
    '“': '"', '”': '"', '…': '...', '•': '-',
    ' ': ' ', '­': '', '→': '->', '←': '<-',
    '×': 'x', '÷': '/', '≥': '>=', '≤': '<=', '≠': '!=',
    '°': 'deg', '±': '+/-', '−': '-', '²': '2', '³': '3',
})
_UNICODE_TO_LATIN1.update(str.maketrans({
    '«': '"', '»': '"', '◦': 'o', '▪': '-', '▫': '-', '·': '.',
    '​': '', '‌': '', '‍': '', '﻿': '',
    ' ': '\n', ' ': '\n',
    '↑': '^', '↓': 'v', '⇒': '=>', '⇐': '<=', '↔': '<->',
    '≈': '~', '≡': '==', '∞': 'inf', '√': 'sqrt', '∑': 'sum',
    '∏': 'prod', '∫': 'int', '∂': 'd',
    'α': 'alpha', 'β': 'beta', 'γ': 'gamma', 'δ': 'delta',
    'ε': 'epsilon', 'ζ': 'zeta', 'η': 'eta', 'θ': 'theta',
    'ι': 'iota', 'κ': 'kappa', 'λ': 'lambda', 'μ': 'mu',
    'ν': 'nu', 'ξ': 'xi', 'π': 'pi', 'ρ': 'rho',
    'σ': 'sigma', 'τ': 'tau', 'υ': 'upsilon', 'φ': 'phi',
    'χ': 'chi', 'ψ': 'psi', 'ω': 'omega',
    'Δ': 'Delta', 'Σ': 'Sigma', 'Π': 'Pi', 'Ω': 'Omega',
    'Γ': 'Gamma', 'Λ': 'Lambda', 'Θ': 'Theta', 'Φ': 'Phi',
    '¹': '1', '⌀': 'diam', '⊥': 'perp', '∥': '||', '∝': 'prop',
}))

_UNICODE_FONT_CANDIDATES = [
    ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
     '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
    ('/usr/share/fonts/dejavu/DejaVuSans.ttf',
     '/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf'),
    ('/usr/share/fonts/TTF/DejaVuSans.ttf',
     '/usr/share/fonts/TTF/DejaVuSans-Bold.ttf'),
    ('C:/Windows/Fonts/arial.ttf', 'C:/Windows/Fonts/arialbd.ttf'),
    ('/System/Library/Fonts/Supplemental/Arial.ttf',
     '/System/Library/Fonts/Supplemental/Arial Bold.ttf'),
    ('/Library/Fonts/Arial.ttf', '/Library/Fonts/Arial Bold.ttf'),
]

_INLINE_MD_RE = re.compile(r'\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|~~(.+?)~~|\*([^*\n]+?)\*')
_STRIKE_RGB = (130, 130, 130)


def _to_latin1(text: str) -> str:
    text = text.translate(_UNICODE_TO_LATIN1)
    return text.encode('latin-1', errors='replace').decode('latin-1')


def _find_unicode_fonts():
    """Return (regular_path, bold_path) for a Unicode TTF, or (None, None)."""
    from pathlib import Path as _P
    for reg, bold in _UNICODE_FONT_CANDIDATES:
        if _P(reg).exists():
            bold_path = bold if _P(bold).exists() else reg
            return reg, bold_path
    return None, None


def _safe_style(font_family: str, style: str) -> str:
    if font_family != 'Helvetica':
        style = style.replace('I', '')
    return style


def _pdf_write_inline(pdf, text: str, body_rgb: tuple, bold_rgb: tuple, line_h: float,
                      font_family: str = 'Helvetica') -> None:
    """Render inline Markdown (**bold**, *italic*, ~~strike~~) as styled spans."""
    def _emit(segment: str, style: str, color: tuple) -> None:
        if not segment:
            return
        pdf.set_font(font_family, _safe_style(font_family, style), 10)
        pdf.set_text_color(*color)
        pdf.write(line_h, segment)

    pos = 0
    for m in _INLINE_MD_RE.finditer(text):
        _emit(text[pos:m.start()], '', body_rgb)
        if m.group(1) is not None:
            _emit(m.group(1), 'BI', bold_rgb)
        elif m.group(2) is not None:
            _emit(m.group(2), 'B', bold_rgb)
        elif m.group(3) is not None:
            _emit(m.group(3), '', _STRIKE_RGB)
        elif m.group(4) is not None:
            _emit(m.group(4), 'I', body_rgb)
        pos = m.end()
    _emit(text[pos:], '', body_rgb)


def _build_pdf_bytes_from_markdown(markdown_text: str, title: str = "Analysis Report") -> bytes:
    """Generate a PDF from LLM markdown using fpdf2. Pure Python, zero system-lib deps."""
    import datetime as _dt
    from fpdf import FPDF

    NAVY    = (47,  84, 150)
    NAVY_D  = (26,  56, 107)
    WHITE   = (255, 255, 255)
    BODY    = (38,  38,  38)
    SECT_BG = (232, 239, 249)
    ROW_ALT = (242, 246, 252)
    CODE_BG = (245, 246, 248)
    RULE    = (196, 213, 235)
    GRAY    = (110, 120, 140)

    LH = 5.5
    M  = 22
    W  = 166

    gen_date = _dt.datetime.now(_dt.timezone.utc).strftime('%d %b %Y')

    unicode_reg, unicode_bold = _find_unicode_fonts()
    if unicode_reg:
        font_family = 'AppFont'
        font_mono   = 'AppFont'
        safe_title  = title
        safe_md     = markdown_text
    else:
        font_family = 'Helvetica'
        font_mono   = 'Courier'
        safe_title  = _to_latin1(title)
        safe_md     = _to_latin1(markdown_text)

    class _PDF(FPDF):
        def footer(self):
            self.set_y(-15)
            self.set_x(M)
            self.set_font(font_family, '', 8)
            self.set_text_color(*GRAY)
            self.cell(W // 2, 8, 'QualiBOT - Document Comparison', align='L')
            self.cell(W // 2, 8, f'Page {self.page_no()}', align='R')

    pdf = _PDF(format='A4')
    pdf.set_margins(M, M, M)
    pdf.set_auto_page_break(auto=True, margin=26)
    if unicode_reg:
        pdf.add_font(font_family, '', unicode_reg)
        pdf.add_font(font_family, 'B', unicode_bold)
    pdf.add_page()

    BANNER_H = 32
    TITLE_W  = W - 14

    pdf.set_fill_color(*NAVY)
    pdf.rect(M, M, W, BANNER_H, 'F')
    pdf.set_fill_color(*NAVY_D)
    pdf.rect(M, M, 5, BANNER_H, 'F')
    pdf.set_fill_color(70, 130, 210)
    pdf.rect(M, M + BANNER_H - 2, W, 2, 'F')

    pdf.set_text_color(*WHITE)
    for fsize in (16, 14, 12, 10):
        pdf.set_font(font_family, 'B', fsize)
        if pdf.get_string_width(safe_title) <= TITLE_W:
            break
    else:
        while len(safe_title) > 4 and pdf.get_string_width(safe_title + '...') > TITLE_W:
            safe_title = safe_title[:-1]
        safe_title += '...'

    line_h = fsize * 0.4
    pdf.set_xy(M + 9, M + (BANNER_H - line_h) / 2 - 3)
    pdf.cell(TITLE_W, line_h + 2, safe_title, align='L')
    pdf.set_font(font_family, '', 8)
    pdf.set_text_color(200, 215, 240)
    pdf.set_xy(M + 9, M + (BANNER_H - line_h) / 2 + line_h - 1)
    pdf.cell(TITLE_W, 5, f'Generated {gen_date}', align='L')

    pdf.set_y(M + BANNER_H + 6)

    def _wrap_text(text: str, max_w: float, style: str = '', size: float = 10) -> list[str]:
        pdf.set_font(font_family, _safe_style(font_family, style), size)

        def _break_long(word: str) -> list[str]:
            if pdf.get_string_width(word) <= max_w:
                return [word]
            chunks: list[str] = []
            piece = ''
            for ch in word:
                if piece and pdf.get_string_width(piece + ch) > max_w:
                    chunks.append(piece)
                    piece = ch
                else:
                    piece += ch
            if piece:
                chunks.append(piece)
            return chunks

        lines: list[str] = []
        cur = ''
        for word in str(text).split():
            for part in _break_long(word):
                test = (cur + ' ' + part).strip()
                if pdf.get_string_width(test) <= max_w:
                    cur = test
                else:
                    if cur:
                        lines.append(cur)
                    cur = part
        if cur:
            lines.append(cur)
        return lines or ['']

    def _wrap(text: str, col_w: float, bold: bool = False) -> list[str]:
        pdf.set_font(font_family, 'B' if bold else '', 8)
        avail = col_w - 3

        def _break_long(word: str) -> list[str]:
            if pdf.get_string_width(word) <= avail:
                return [word]
            chunks: list[str] = []
            piece = ''
            for ch in word:
                if piece and pdf.get_string_width(piece + ch) > avail:
                    chunks.append(piece)
                    piece = ch
                else:
                    piece += ch
            if piece:
                chunks.append(piece)
            return chunks

        lines: list[str] = []
        cur = ''
        for word in str(text).split():
            for part in _break_long(word):
                test = (cur + ' ' + part).strip()
                if pdf.get_string_width(test) <= avail:
                    cur = test
                else:
                    if cur:
                        lines.append(cur)
                    cur = part
        if cur:
            lines.append(cur)
        return lines or ['']

    def _flush_table(table_lines: list[str]) -> None:
        T_LH  = 4.5
        T_PAD = 1.5

        parsed: list[list[str]] = []
        for raw_line in table_lines:
            s = raw_line.strip()
            if not s:
                continue
            cells = [c.strip() for c in s.strip('|').split('|')]
            if all(re.match(r'^[-:]+$', c or '-') for c in cells):
                continue
            parsed.append(cells)

        if not parsed:
            return

        n_cols = max(1, max(len(r) for r in parsed))

        def _pad(row: list[str]) -> list[str]:
            return (list(row) + [''] * n_cols)[:n_cols]

        parsed = [_pad(r) for r in parsed]
        headers   = parsed[0]
        data_rows = parsed[1:]
        col_w = W / n_cols

        def _draw_row(cells: list[str], is_header: bool, row_idx: int = 0) -> None:
            wrapped = [_wrap(c, col_w, bold=is_header) for c in cells]
            n_lines = max(len(w) for w in wrapped)
            row_h   = n_lines * T_LH + T_PAD * 2

            if pdf.get_y() + row_h > pdf.h - 28:
                pdf.add_page()
                if not is_header:
                    _draw_row(headers, is_header=True)

            y = pdf.get_y()
            pdf.set_draw_color(*RULE)
            pdf.set_line_width(0.2)

            for ci, (wlines, _) in enumerate(zip(wrapped, cells)):
                x = M + ci * col_w
                if is_header:
                    pdf.set_fill_color(*NAVY)
                    pdf.set_text_color(*WHITE)
                    pdf.set_font(font_family, 'B', 8)
                else:
                    pdf.set_fill_color(*ROW_ALT if row_idx % 2 == 0 else WHITE)
                    pdf.set_text_color(*BODY)
                    pdf.set_font(font_family, '', 8)

                pdf.rect(x, y, col_w, row_h, 'DF')

                ty = y + T_PAD
                for wline in wlines:
                    pdf.set_xy(x + 1.5, ty)
                    pdf.cell(col_w - 3, T_LH, wline,
                             align='C' if is_header else 'L')
                    ty += T_LH

            pdf.set_y(y + row_h)

        pdf.ln(2)
        _draw_row(headers, is_header=True)
        for ri, row in enumerate(data_rows):
            _draw_row(row, is_header=False, row_idx=ri)
        pdf.ln(4)

    in_code    = False
    prev_blank = True
    table_buf: list[str] = []

    for raw in safe_md.splitlines():
        line = raw.rstrip()

        if line.startswith('|'):
            table_buf.append(line)
            prev_blank = False
            continue

        if table_buf:
            _flush_table(table_buf)
            table_buf = []
            prev_blank = False

        if line.startswith('```'):
            in_code = not in_code
            if in_code:
                pdf.ln(2)
            continue
        if in_code:
            y = pdf.get_y()
            pdf.set_fill_color(*CODE_BG)
            pdf.rect(M, y, W, LH + 1, 'F')
            pdf.set_font(font_mono, '', 8.5)
            pdf.set_text_color(*BODY)
            pdf.set_x(M + 3)
            pdf.cell(W - 3, LH + 1, line, new_x='LMARGIN', new_y='NEXT')
            continue

        if not line.strip():
            pdf.ln(3)
            prev_blank = True
            continue

        if re.match(r'^## [^#]', line):
            text = line[3:].strip()
            pdf.ln(3 if prev_blank else 6)
            hlines = _wrap_text(text, W - 11, 'B', 11)
            box_h = max(11, len(hlines) * 6 + 4)
            if pdf.get_y() + box_h > pdf.h - 26:
                pdf.add_page()
            y = pdf.get_y()
            pdf.set_fill_color(*SECT_BG)
            pdf.rect(M, y, W, box_h, 'F')
            pdf.set_fill_color(*NAVY)
            pdf.rect(M, y, 4, box_h, 'F')
            pdf.set_text_color(*NAVY_D)
            pdf.set_font(font_family, 'B', 11)
            ty = y + (box_h - len(hlines) * 6) / 2
            for hl in hlines:
                pdf.set_xy(M + 7, ty)
                pdf.cell(W - 9, 6, hl, align='L')
                ty += 6
            pdf.set_y(y + box_h + 3)
            prev_blank = False
            continue

        if re.match(r'^### [^#]', line):
            text = line[4:].strip()
            pdf.ln(2 if prev_blank else 5)
            hlines = _wrap_text(text, W, 'B', 10)
            pdf.set_text_color(*NAVY)
            pdf.set_font(font_family, 'B', 10)
            for hl in hlines:
                pdf.set_x(M)
                pdf.cell(W, 6, hl, align='L', new_x='LMARGIN', new_y='NEXT')
            y = pdf.get_y()
            pdf.set_draw_color(*RULE)
            pdf.set_line_width(0.3)
            pdf.line(M, y + 1, M + W, y + 1)
            pdf.set_line_width(0.2)
            pdf.ln(2)
            prev_blank = False
            continue

        if re.match(r'^# [^#]', line):
            text = line[2:].strip()
            pdf.ln(3 if prev_blank else 6)
            hlines = _wrap_text(text, W, 'B', 13)
            pdf.set_text_color(*NAVY)
            pdf.set_font(font_family, 'B', 13)
            for hl in hlines:
                pdf.set_x(M)
                pdf.cell(W, 7, hl, new_x='LMARGIN', new_y='NEXT')
            y = pdf.get_y()
            pdf.set_draw_color(*NAVY)
            pdf.set_line_width(0.6)
            pdf.line(M, y, M + W, y)
            pdf.set_line_width(0.2)
            pdf.ln(4)
            prev_blank = False
            continue

        m = re.match(r'^(\s*)([-*+])\s+(.*)', line)
        if m:
            indent = len(m.group(1)) // 2
            text   = m.group(3)
            bx     = M + indent * 5
            tx     = bx + 5
            y = pdf.get_y()
            pdf.set_fill_color(*NAVY)
            pdf.rect(bx + 0.5, y + 2.5, 2.2, 2.2, 'F')
            pdf.set_xy(tx, y)
            _pdf_write_inline(pdf, text, BODY, NAVY_D, LH, font_family=font_family)
            pdf.ln(LH + 1.5)
            prev_blank = False
            continue

        m = re.match(r'^(\s*)(\d+)[.)]\s+(.*)', line)
        if m:
            indent = len(m.group(1)) // 2
            num    = m.group(2)
            text   = m.group(3)
            nx     = M + indent * 5
            tx     = nx + 8
            y = pdf.get_y()
            pdf.set_font(font_family, 'B', 10)
            pdf.set_text_color(*NAVY)
            pdf.set_xy(nx, y)
            pdf.cell(7, LH, f'{num}.')
            pdf.set_xy(tx, y)
            _pdf_write_inline(pdf, text, BODY, NAVY_D, LH, font_family=font_family)
            pdf.ln(LH + 1.5)
            prev_blank = False
            continue

        if re.match(r'^\*\*[^*].+\*\*$', line):
            inner = line[2:-2]
            pdf.ln(2)
            blines = _wrap_text(inner, W - 8, 'B', 10)
            box_h = max(9, len(blines) * 5.5 + 3)
            if pdf.get_y() + box_h > pdf.h - 26:
                pdf.add_page()
            y = pdf.get_y()
            pdf.set_fill_color(*SECT_BG)
            pdf.rect(M, y, W, box_h, 'F')
            pdf.set_text_color(*NAVY_D)
            pdf.set_font(font_family, 'B', 10)
            ty = y + (box_h - len(blines) * 5.5) / 2
            for bl in blines:
                pdf.set_xy(M + 4, ty)
                pdf.cell(W - 8, 5.5, bl, align='L')
                ty += 5.5
            pdf.set_y(y + box_h + 3)
            prev_blank = False
            continue

        if re.match(r'^[-*_]{3,}$', line.strip()):
            pdf.ln(2)
            y = pdf.get_y()
            pdf.set_draw_color(*RULE)
            pdf.set_line_width(0.4)
            pdf.line(M, y, M + W, y)
            pdf.set_line_width(0.2)
            pdf.ln(4)
            prev_blank = True
            continue

        pdf.set_x(M)
        _pdf_write_inline(pdf, line, BODY, NAVY_D, LH, font_family=font_family)
        pdf.ln(LH + 1)
        prev_blank = False

    if table_buf:
        _flush_table(table_buf)

    return bytes(pdf.output())
