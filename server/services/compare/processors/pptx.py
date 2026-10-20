"""PPTX / PPT processor — three processing methods.

standard:    paragraph semantic diff on slide text + focused Markdown bullet-list output.
structured:  paragraph semantic diff on slide text + strict JSON array output (for Excel export).
comparative: section canonical diff on slide text + section-grouped output.

Embedded images go through the same perceptual hashing as PDF/DOCX (_compute_image_diff_pairs in _diff_engines),
so they get the same MODIFIED / REMOVED / ADDED / ORIENTATION categories.
"""

import base64
import io
import logging
from typing import Any, Dict, List, Tuple

from server.services.compare.processors.text_diff import Images, TextDiffProcessor
from server.services.compare.processors._diff_engines import (
    dhash,
    hash16,
    images_are_similar,
)

logger = logging.getLogger(__name__)

_IMG_MAX_PX = 1024
_IMG_JPEG_QUALITY = 80


def _shapes_text(shapes) -> List[str]:
    """Recursively extract text from shapes, including tables and groups.

    :param shapes: shapes of a slide or of a group
    :return: the text lines: paragraphs, table rows joined by ``|``, and ``[Chart]`` lines
    """
    texts: List[str] = []
    for shape in shapes:
        try:
            if shape.shape_type == 6:  # group
                texts.extend(_shapes_text(shape.shapes))
                continue
        except Exception:
            pass
        try:
            # Charts carry technical content (title, series, categories) that
            # would otherwise be invisible to the text diff.
            if getattr(shape, 'has_chart', False):
                chart = shape.chart
                parts: List[str] = []
                try:
                    if chart.has_title:
                        title = chart.chart_title.text_frame.text.strip()
                        if title:
                            parts.append(title)
                except Exception:
                    pass
                try:
                    names = [str(s.name) for s in chart.series if getattr(s, 'name', None)]
                    if names:
                        parts.append('series: ' + ', '.join(names))
                except Exception:
                    pass
                try:
                    cats = [str(c) for c in chart.plots[0].categories if c is not None]
                    if cats:
                        parts.append('categories: ' + ', '.join(cats[:20]))
                except Exception:
                    pass
                if parts:
                    texts.append('[Chart] ' + ' | '.join(parts))
                continue
        except Exception:
            pass
        try:
            if shape.has_table:
                for row in shape.table.rows:
                    row_cells = []
                    for cell in row.cells:
                        try:
                            t = cell.text_frame.text.strip()
                            if t:
                                row_cells.append(t)
                        except Exception:
                            pass
                    if row_cells:
                        texts.append(' | '.join(row_cells))
                continue
        except Exception:
            pass
        try:
            if hasattr(shape, 'has_text_frame') and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        texts.append(text)
        except Exception:
            pass
    return texts


def _collect_shape_images(shapes, slide_num: int, images: Dict[str, Dict[str, Any]]) -> None:
    """Recursively extract images from shapes, dedupe by perceptual dhash, record slides.

    Builds the same dict shape used by DOCX/PDF processors:
        { dhash_str: { 'b64': str, 'pages': [slide_nums] } }
    so the shared pairing engine in _diff_engines can be applied consistently.

    :param shapes: shapes of a slide or of a group
    :param slide_num: number of the slide, from 1
    :param images: dict filled in place, keyed by perceptual hash
    """
    from PIL import Image as PILImage
    for shape in shapes:
        try:
            if shape.shape_type == 6:  # group
                _collect_shape_images(shape.shapes, slide_num, images)
                continue
        except Exception:
            pass
        try:
            if shape.shape_type not in (13, 14):  # picture, placeholder picture
                continue
            img_bytes = shape.image.blob
            # Normalise to JPEG with consistent size and quality so the hash
            # remains stable across small re-saves of the same source image.
            img = PILImage.open(io.BytesIO(img_bytes)).convert('RGB')
            w, h = img.size
            scale = min(1.0, _IMG_MAX_PX / max(w, h, 1))
            if scale < 1.0:
                img = img.resize(
                    (max(1, int(w * scale)), max(1, int(h * scale))),
                    PILImage.Resampling.LANCZOS,
                )
            buf = io.BytesIO()
            img.save(buf, format='JPEG', quality=_IMG_JPEG_QUALITY)
            jpeg_bytes = buf.getvalue()

            h_key = dhash(jpeg_bytes)
            # Near-duplicate pictures (one slide reusing another's picture, re-compressed) within a small Hamming budget collapse
            # into one entry listing several slides.
            matched = next((k for k in images if images_are_similar(k, h_key)), None)
            if matched:
                if slide_num not in images[matched]['pages']:
                    images[matched]['pages'].append(slide_num)
            else:
                b64 = base64.b64encode(jpeg_bytes).decode('ascii')
                images[h_key] = {'b64': b64, 'pages': [slide_num], 'hash16': hash16(jpeg_bytes)}
        except Exception as e:
            logger.debug('PPTX image extraction skipped (slide %d): %s', slide_num, e)


def _extract_pptx(pptx_bytes: bytes) -> Tuple[List[str], Dict[str, Dict[str, Any]]]:
    """Return (slide_text_lines, images_dict).

    Each text line is tagged with a trailing [Slide N] tag so the diff engine
    can track which slide a change originates from. Images are returned in the
    same shape used by DOCX/PDF (see _collect_shape_images).

    :param pptx_bytes: content of the .pptx file
    :return: ``(slide text lines, images keyed by hash)``
    """
    from pptx import Presentation
    prs = Presentation(io.BytesIO(pptx_bytes))
    slide_texts: List[str] = []
    images: Dict[str, Dict[str, Any]] = {}
    for slide_num, slide in enumerate(prs.slides, start=1):
        lines = [f'{t} [Slide {slide_num}]' for t in _shapes_text(slide.shapes)]
        slide_texts.append('\n'.join(lines) if lines else f'[Slide {slide_num}]')
        _collect_shape_images(slide.shapes, slide_num, images)
    return slide_texts, images


# --- Processor ---

class PptxProcessor(TextDiffProcessor):
    file_type = 'pptx'

    page_label = 'Slide'
    warn_comparative = False

    def extract(self, old_bytes, old_name, new_bytes, new_name) -> Tuple[str, str, Images, Images]:
        """Slide text and images of the two presentations.

        :param old_bytes: content of the old presentation
        :param old_name: file name of the old presentation
        :param new_bytes: content of the new presentation
        :param new_name: file name of the new presentation
        :return: ``(old text, new text, old images, new images)``
        """
        old_slides, old_imgs = _extract_pptx(old_bytes)
        new_slides, new_imgs = _extract_pptx(new_bytes)
        return '\n'.join(old_slides), '\n'.join(new_slides), old_imgs, new_imgs

    def paragraph_notes(self, filtered) -> str:
        """The note that opens the paragraph diff of a presentation.

        :param filtered: number of trivial entries the diff filtered out
        :return: the note
        """
        return (f'Global paragraph alignment across slides: {filtered} trivial lines filtered.\n'
                'Similar text blocks shown as MODIFIED with ~~removed~~ and **added** words inline.\n'
                'Each block tagged with its source slide.')

    def comparative_notes(self, filtered) -> str:
        """The note that opens the section diff of a presentation.

        :param filtered: number of trivial entries the diff filtered out
        :return: the note
        """
        return f'Section canonical diff across slides: {filtered} trivial pairs filtered.'
